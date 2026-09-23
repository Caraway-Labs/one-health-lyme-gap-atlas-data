from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from jsonschema import Draft202012Validator
from openpyxl import Workbook

from lyme_gap_atlas_data import tick_surveillance as tick
from lyme_gap_atlas_data.artifacts import create_artifact
from lyme_gap_atlas_data.migrations import (
    DEV_DATABASE,
    load_migrations,
    migration_execution_role,
    migration_plan,
    render_migration,
)
from lyme_gap_atlas_data.tick_contract import canonical_observation_id
from lyme_gap_atlas_data.tick_normalization import convert_value, normalize_value


def workbook_bytes(
    *,
    headers: tuple[str, ...] = tick.EXPECTED_HEADERS,
    rows: list[tuple[object, ...]] | None = None,
) -> bytes:
    workbook = Workbook()
    agreement = workbook.active
    agreement.title = "Data Use Agreement"
    agreement.append(["Access to ArboNET Tick Module data is limited to the Requestor."])
    agreement.append(["These data should not be provided to other persons."])
    agreement.append(["ArboNET will be appropriately referenced."])
    agreement.append(["A final copy of publications will be provided to CDC."])
    agreement.append(["ArboNET is a passive surveillance system."])
    terms = workbook.create_sheet("Classification Terms")
    terms.append(["County Classification", "Definition"])
    terms.append(["Established", "Reviewed publisher definition"])
    terms.append(["Reported", "Reviewed publisher definition"])
    terms.append(["No records", "No records should not be interpreted as ticks being absent."])
    data = workbook.create_sheet("Ixodes records 2025")
    data.append(["Ixodes status through Dec. 31, 2025"])
    data.append(list(headers))
    source_rows = rows or [
        (
            f"{1000 + index:05d}",
            "Fixture State",
            f"Fixture County {index}",
            "Established" if index % 2 else "No records",
            "Fixture citation",
            "Reported",
            "Fixture citation",
        )
        for index in range(1, 31)
    ]
    for row in source_rows:
        data.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def profile() -> dict[str, object]:
    return tick.load_tick_profile()


def write_evidence_bundle(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    landing = (
        b"Tick Surveillance Data Sets No records Established "
        b"Public_Use_Ixodes_County_Table_2026_03252026.xlsx"
    )
    workbook = workbook_bytes()
    (tmp_path / "landing.html").write_bytes(landing)
    (tmp_path / "workbook.xlsx").write_bytes(workbook)
    base_digest = "sha256:" + "a" * 64
    envelope_digest = "sha256:" + "b" * 64
    manifest = {
        "manifest_version": 1,
        "acquisition_route": tick.EVIDENCE_ROUTE,
        "retrieved_at": tick.datetime.now(tick.UTC).isoformat(),
        "github": {
            "repository": "Caraway-Labs/one-health-lyme-gap-atlas-data",
            "run_id": "123456789",
            "run_attempt": "1",
            "sha": "c" * 40,
        },
        "base_image_digest": base_digest,
        "resources": [
            {
                "purpose": "SOURCE_LANDING_PAGE",
                "filename": "landing.html",
                "requested_url": profile()["landing_page_url"],
                "final_url": profile()["landing_page_url"],
                "status_code": 200,
                "media_type": "text/html",
                "byte_count": len(landing),
                "sha256": hashlib.sha256(landing).hexdigest(),
                "etag": None,
                "last_modified": None,
            },
            {
                "purpose": "SOURCE_WORKBOOK_EVIDENCE",
                "filename": "workbook.xlsx",
                "requested_url": profile()["endpoint_template"],
                "final_url": profile()["endpoint_template"],
                "status_code": 200,
                "media_type": tick.XLSX_MEDIA_TYPE,
                "byte_count": len(workbook),
                "sha256": hashlib.sha256(workbook).hexdigest(),
                "etag": None,
                "last_modified": None,
            },
        ],
    }
    (tmp_path / "acquisition-manifest.json").write_text(json.dumps(manifest))
    monkeypatch.setenv("TICK_EVIDENCE_BASE_IMAGE_DIGEST", base_digest)
    monkeypatch.setenv("TICK_EVIDENCE_ENVELOPE_DIGEST", envelope_digest)
    monkeypatch.setenv("TICK_EVIDENCE_GITHUB_RUN_ID", "123456789")
    return tmp_path


def write_operator_evidence_bundle(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    landing = b"%PDF-1.7\nfixture operator print\n%%EOF\n"
    workbook = workbook_bytes()
    (tmp_path / "landing.pdf").write_bytes(landing)
    (tmp_path / "workbook.xlsx").write_bytes(workbook)
    base_digest = "sha256:" + "a" * 64
    envelope_digest = "sha256:" + "b" * 64
    retrieval_id = "12345678-1234-4234-8234-123456789abc"
    retrieved_at = tick.datetime.now(tick.UTC).isoformat()
    manifest = {
        "manifest_version": 2,
        "acquisition_route": tick.OPERATOR_EVIDENCE_ROUTE,
        "retrieved_at": retrieved_at,
        "operator": {
            "retrieval_id": retrieval_id,
            "acquisition_method": "BROWSER_DOWNLOAD_AND_PRINT",
            "attestation": "FILES_SAVED_FROM_PINNED_FIRST_PARTY_CDC_PAGE",
        },
        "base_image_digest": base_digest,
        "resources": [
            {
                "purpose": "SOURCE_LANDING_PAGE_PRINT",
                "filename": "landing.pdf",
                "requested_url": profile()["landing_page_url"],
                "final_url": profile()["landing_page_url"],
                "status_code": None,
                "http_status_observed": False,
                "media_type": tick.PDF_MEDIA_TYPE,
                "byte_count": len(landing),
                "sha256": hashlib.sha256(landing).hexdigest(),
                "etag": None,
                "last_modified": None,
                "transport": "BROWSER_PRINT_TO_PDF",
                "source_file_modified_at": retrieved_at,
            },
            {
                "purpose": "SOURCE_WORKBOOK_EVIDENCE",
                "filename": "workbook.xlsx",
                "requested_url": profile()["endpoint_template"],
                "final_url": profile()["endpoint_template"],
                "status_code": None,
                "http_status_observed": False,
                "media_type": tick.XLSX_MEDIA_TYPE,
                "byte_count": len(workbook),
                "sha256": hashlib.sha256(workbook).hexdigest(),
                "etag": None,
                "last_modified": None,
                "transport": "BROWSER_DOWNLOAD",
                "source_file_modified_at": retrieved_at,
            },
        ],
    }
    (tmp_path / "acquisition-manifest.json").write_text(json.dumps(manifest))
    monkeypatch.setenv("TICK_EVIDENCE_BASE_IMAGE_DIGEST", base_digest)
    monkeypatch.setenv("TICK_EVIDENCE_ENVELOPE_DIGEST", envelope_digest)
    monkeypatch.setenv("TICK_EVIDENCE_OPERATOR_RETRIEVAL_ID", retrieval_id)
    return tmp_path


def test_canonical_tick_contract_has_required_semantics_and_examples() -> None:
    schema = json.loads(
        Path(
            "docs/contracts/tick-surveillance/canonical-tick-surveillance-v1.schema.json"
        ).read_text(encoding="utf-8")
    )
    assert schema["properties"]["county_fips"]["pattern"] == "^[0-9]{5}$"
    assert "county_fips" not in schema["required"]
    assert schema["properties"]["method_version"]["enum"] == [
        "tick-surveillance-v1",
        "tick-surveillance-v1.1",
        "tick-surveillance-v1.2",
    ]
    assert set(schema["properties"]["observation_type"]["enum"]) == {
        "VECTOR_PRESENCE_STATUS",
        "COLLECTION_ABUNDANCE",
        "PATHOGEN_PRESENCE_STATUS",
        "PATHOGEN_TESTING",
    }
    missingness = set(schema["properties"]["missingness"]["additionalProperties"]["enum"])
    assert {"NULL", "UNKNOWN", "SUPPRESSED", "NOT_REPORTED", "NOT_APPLICABLE"} <= missingness
    contract = Path("docs/contracts/tick-surveillance/canonical-tick-surveillance-v1.md").read_text(
        encoding="utf-8"
    )
    assert contract.count('"canonical_observation_id"') == 3
    assert "NO_RECORDS" in contract
    assert "not evidence that ticks or pathogens are absent" in contract
    assert "NOT_COUNTY_REPRESENTATIVE" in contract


def test_governed_normalization_registry_is_independently_schema_valid() -> None:
    registry_path = Path("docs/contracts/tick-surveillance/tick-surveillance-normalization-v1.json")
    schema_path = Path(
        "docs/contracts/tick-surveillance/tick-surveillance-normalization-v1.schema.json"
    )
    registry = json.loads(registry_path.read_text(encoding="utf-8"))
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    assert list(Draft202012Validator(schema).iter_errors(registry)) == []
    assert registry["registry_version"] == "1.0.4"
    assert "comparability" in " ".join(registry["scope"]["non_goals"]).lower()


def test_governed_taxon_mappings_preserve_source_value_and_pin_version() -> None:
    cdc = normalize_value(
        field="tick_taxon",
        source_value="Ixodes_scapularis",
        publisher="CDC ArboNET Tick Module",
        dataset_id="cdc-ixodes-county-status-2025",
        source_version="2025",
    )
    assert cdc.status == "APPROVED"
    assert cdc.canonical_id == "IXODES_SCAPULARIS"
    assert cdc.canonical_label == "Ixodes scapularis"
    assert cdc.as_contract_value()["source_value"] == "Ixodes_scapularis"
    assert cdc.mapping_rule_id == "TAXON_CDC_SCAPULARIS_V1"
    assert cdc.registry_version == "1.0.4"

    aggregate = normalize_value(
        field="tick_taxon",
        source_value="Ixodes scapularis or Ixodes pacificus",
        publisher="CDC ArboNET Tick Module",
        dataset_id="cdc-ixodes-pathogen-status-2025",
        source_version="2025",
    )
    assert aggregate.canonical_id == "IXODES_SCAPULARIS_OR_PACIFICUS"
    assert aggregate.canonical_id != cdc.canonical_id


def test_neon_amblyomma_mapping_is_exact_and_source_context_specific() -> None:
    context = {
        "publisher": "NSF NEON",
        "dataset_id": "DP1.10093.001",
        "source_version": "RELEASE-2026",
    }
    approved = normalize_value(field="tick_taxon", source_value="Amblyomma americanum", **context)
    altered_case = normalize_value(
        field="tick_taxon", source_value="Amblyomma Americanum", **context
    )
    wrong_context = normalize_value(
        field="tick_taxon",
        source_value="Amblyomma americanum",
        publisher="NSF NEON",
        dataset_id="DP1.10092.001",
        source_version="RELEASE-2026",
    )
    ixodes = normalize_value(field="tick_taxon", source_value="Ixodes scapularis", **context)

    assert (approved.status, approved.canonical_id, approved.mapping_rule_id) == (
        "APPROVED",
        "AMBLYOMMA_AMERICANUM",
        "TAXON_NEON_AMERICANUM_V1",
    )
    assert (altered_case.status, altered_case.canonical_id, altered_case.mapping_rule_id) == (
        "UNKNOWN",
        None,
        None,
    )
    assert (wrong_context.status, wrong_context.canonical_id) == ("UNKNOWN", None)
    assert (ixodes.status, ixodes.canonical_id, ixodes.mapping_rule_id) == (
        "APPROVED",
        "IXODES_SCAPULARIS",
        "TAXON_NEON_SCAPULARIS_V1",
    )
def test_neon_life_stage_mappings_are_exact_and_preserve_approved_casing() -> None:
    context = {
        "publisher": "NSF NEON",
        "dataset_id": "DP1.10093.001",
        "source_version": "RELEASE-2026",
    }
    lower = normalize_value(field="life_stage", source_value="nymph", **context)
    capitalized = normalize_value(field="life_stage", source_value="Nymph", **context)
    unapproved_case = normalize_value(field="life_stage", source_value="NYMPH", **context)

    assert (lower.status, lower.canonical_id, lower.mapping_rule_id) == (
        "APPROVED",
        "NYMPH",
        "LIFE_STAGE_NEON_NYMPH_V1",
    )
    assert (capitalized.status, capitalized.canonical_id, capitalized.mapping_rule_id) == (
        "APPROVED",
        "NYMPH",
        "LIFE_STAGE_NEON_NYMPH_CAPITALIZED_V1",
    )
    assert (
        unapproved_case.status,
        unapproved_case.canonical_id,
        unapproved_case.mapping_rule_id,
    ) == (
        "UNKNOWN",
        None,
        None,
    )


def test_governed_normalization_fails_closed_for_unknown_taxa_and_pathogens() -> None:
    unknown = normalize_value(
        field="tick_taxon",
        source_value="Ixodes inventedus",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    assert (unknown.status, unknown.canonical_id, unknown.mapping_rule_id) == (
        "UNKNOWN",
        None,
        None,
    )
    pathogen = normalize_value(
        field="pathogen_target",
        source_value="Borrelia_mayonii",
        publisher="CDC ArboNET Tick Module",
        dataset_id="cdc-ixodes-pathogen-status-2025",
        source_version="2025",
    )
    assert pathogen.status == "APPROVED"
    assert pathogen.canonical_id == "BORRELIA_MAYONII"
    assert pathogen.mapping_rule_id == "PATHOGEN_CDC_BMAYONII_V1"

    unknown_pathogen = normalize_value(
        field="pathogen_target",
        source_value="Unknown pathogen target",
        publisher="NSF NEON",
        dataset_id="DP1.10092.001",
        source_version="RELEASE-2026",
    )
    assert (unknown_pathogen.status, unknown_pathogen.canonical_id) == ("UNKNOWN", None)


def test_neon_pathogen_mappings_are_exact_and_source_context_specific() -> None:
    context = {
        "publisher": "NSF NEON",
        "dataset_id": "DP1.10092.001",
        "source_version": "RELEASE-2026",
    }
    expected = {
        "Anaplasma phagocytophilum": "ANAPLASMA_PHAGOCYTOPHILUM",
        "Babesia microti": "BABESIA_MICROTI",
        "Borrelia mayonii": "BORRELIA_MAYONII",
        "Borrelia miyamotoi": "BORRELIA_MIYAMOTOI",
        "Ehrlichia muris-like": "EHRLICHIA_MURIS_LIKE_AGENT",
        "Borrelia sp.": "BORRELIA_SP",
    }

    results = {
        source_value: normalize_value(field="pathogen_target", source_value=source_value, **context)
        for source_value in expected
    }

    assert {value: result.canonical_id for value, result in results.items()} == expected
    assert results["Ehrlichia muris-like"].mapping_rule_id == (
        "PATHOGEN_NEON_EHRLICHIA_MURIS_LIKE_V1"
    )
    assert results["Borrelia sp."].canonical_id not in {
        "BORRELIA_BURGDORFERI_SENSU_LATO",
        "BORRELIA_MAYONII",
        "BORRELIA_MIYAMOTOI",
    }

    altered = normalize_value(field="pathogen_target", source_value="Borrelia Sp.", **context)
    wrong_context = normalize_value(
        field="pathogen_target",
        source_value="Borrelia sp.",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    assert (altered.status, altered.canonical_id) == ("UNKNOWN", None)
    assert (wrong_context.status, wrong_context.canonical_id) == ("UNKNOWN", None)


def test_neon_non_pathogen_assays_are_retained_as_governed_dispositions() -> None:
    context = {
        "publisher": "NSF NEON",
        "dataset_id": "DP1.10092.001",
        "source_version": "RELEASE-2026",
    }
    qc = normalize_value(field="pathogen_target", source_value="HardTick DNA Quality", **context)
    identification = normalize_value(
        field="pathogen_target", source_value="Ixodes pacificus", **context
    )

    assert (qc.status, qc.canonical_id, qc.disposition) == (
        "UNSUPPORTED",
        None,
        "NON_PATHOGEN_ASSAY_QC",
    )
    assert (identification.status, identification.canonical_id, identification.disposition) == (
        "UNSUPPORTED",
        None,
        "NON_PATHOGEN_TICK_IDENTIFICATION",
    )
    assert qc.as_contract_value()["disposition"] == "NON_PATHOGEN_ASSAY_QC"
    assert identification.as_contract_value()["source_value"] == "Ixodes pacificus"


def test_neon_test_result_mappings_are_exact_and_source_context_specific() -> None:
    context = {
        "publisher": "NSF NEON",
        "dataset_id": "DP1.10092.001",
        "source_version": "RELEASE-2026",
    }
    expected = {
        "positive": ("DETECTED", "RESULT_NEON_POSITIVE_V1"),
        "negative": ("NOT_DETECTED", "RESULT_NEON_NEGATIVE_V1"),
        "Positive": ("DETECTED", "RESULT_NEON_POSITIVE_CAPITALIZED_V1"),
        "Negative": ("NOT_DETECTED", "RESULT_NEON_NEGATIVE_CAPITALIZED_V1"),
    }

    for source_value, (canonical_id, rule_id) in expected.items():
        result = normalize_value(field="test_result", source_value=source_value, **context)
        assert (result.status, result.canonical_id, result.mapping_rule_id) == (
            "APPROVED",
            canonical_id,
            rule_id,
        )

    for source_value in ("POSITIVE", "NEGATIVE", "Positive "):
        result = normalize_value(field="test_result", source_value=source_value, **context)
        assert (result.status, result.canonical_id, result.mapping_rule_id) == (
            "UNKNOWN",
            None,
            None,
        )

    wrong_context = normalize_value(
        field="test_result",
        source_value="Positive",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    assert (wrong_context.status, wrong_context.canonical_id) == ("UNKNOWN", None)


def test_method_vocabulary_normalizes_lexically_without_comparability_decision() -> None:
    drag = normalize_value(
        field="collection_method",
        source_value="drag",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    flag = normalize_value(
        field="collection_method",
        source_value="flag",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    assert (drag.canonical_id, flag.canonical_id) == ("DRAG_CLOTH", "FLAG_CLOTH")
    assert drag.canonical_id != flag.canonical_id
    assert "comparability" not in drag.as_contract_value()


def test_neon_quality_flag_mapping_preserves_the_documented_source_code() -> None:
    quality = normalize_value(
        field="quality_flag",
        source_value="legacyData",
        publisher="NSF NEON",
        dataset_id="DP1.10093.001",
        source_version="RELEASE-2026",
    )
    assert quality.status == "APPROVED"
    assert quality.canonical_id == "LEGACY_DATA"
    assert quality.as_contract_value()["source_value"] == "legacyData"


def test_supported_conversions_require_a_documented_denominator_when_needed() -> None:
    assert convert_value(
        field="effort_unit",
        value=10_000,
        from_canonical_id="SQUARE_METRE",
        to_canonical_id="HECTARE",
        denominator_present=False,
    ) == (1.0, "EFFORT_SQUARE_METRE_TO_HECTARE_V1")
    assert convert_value(
        field="abundance_unit",
        value=2.0,
        from_canonical_id="TICKS_PER_SQUARE_METRE",
        to_canonical_id="TICKS_PER_HECTARE",
        denominator_present=True,
    ) == (20_000.0, "ABUNDANCE_SQUARE_METRE_TO_HECTARE_V1")
    with pytest.raises(ValueError, match="documented denominator"):
        convert_value(
            field="abundance_unit",
            value=2.0,
            from_canonical_id="TICKS_PER_SQUARE_METRE",
            to_canonical_id="TICKS_PER_HECTARE",
            denominator_present=False,
        )
    with pytest.raises(ValueError, match="unsupported"):
        convert_value(
            field="effort_unit",
            value=1.0,
            from_canonical_id="HECTARE",
            to_canonical_id="SQUARE_METRE",
            denominator_present=True,
        )


def test_normalization_contract_envelope_retains_mapping_provenance() -> None:
    result = normalize_value(
        field="test_result",
        source_value="positive",
        publisher="NSF NEON",
        dataset_id="DP1.10092.001",
        source_version="RELEASE-2026",
    )
    record = {
        "canonical_observation_id": "fixture-normalized-test",
        "observation_type": "PATHOGEN_TESTING",
        "tick_species": "Ixodes scapularis",
        "pathogen_name": "Borrelia burgdorferi sensu stricto",
        "ticks_tested": 1,
        "ticks_positive": 1,
        "source_agency": "NSF NEON",
        "source_dataset_id": "DP1.10092.001",
        "source_record_id": "fixture-source-row",
        "data_source_version_id": "RELEASE-2026",
        "ingestion_run_id": "fixture-run",
        "artifact_id": "fixture-artifact",
        "retrieved_at": "2026-09-23T00:00:00Z",
        "method_version": "tick-surveillance-v1.2",
        "reported_or_derived": "HARMONIZED",
        "quality_flags": [],
        "normalization": {
            "registry_id": result.registry_id,
            "registry_version": result.registry_version,
            "mappings": {"test_result": result.as_contract_value()},
        },
    }
    schema = json.loads(
        Path(
            "docs/contracts/tick-surveillance/canonical-tick-surveillance-v1.schema.json"
        ).read_text(encoding="utf-8")
    )
    assert list(Draft202012Validator(schema).iter_errors(record)) == []
    assert record["normalization"]["mappings"]["test_result"]["source_value"] == "positive"


def test_site_event_contract_fixtures_preserve_geography_and_lineage() -> None:
    schema = json.loads(
        Path(
            "docs/contracts/tick-surveillance/canonical-tick-surveillance-v1.schema.json"
        ).read_text(encoding="utf-8")
    )
    fixtures = json.loads(
        Path("tests/fixtures/tick_surveillance/site-event-observations.json").read_text(
            encoding="utf-8"
        )
    )
    validator = Draft202012Validator(schema)
    records = {fixture["case"]: fixture["record"] for fixture in fixtures}
    assert set(records) == {
        "county_native_status",
        "mapped_site_event",
        "unmapped_source_only_site",
        "ambiguous_site_mapping",
        "repeated_event_same_site",
        "multiple_replicates",
    }
    for record in records.values():
        assert list(validator.iter_errors(record)) == []

    county_status = records["county_native_status"]
    assert county_status["method_version"] == "tick-surveillance-v1"
    assert county_status["county_fips"] == "01001"

    mapped = records["mapped_site_event"]
    assert mapped["county_relationship"]["mapping_status"] == "ATLAS_DERIVED_MATCH"
    assert mapped["county_relationship"]["representativeness"] == "NOT_COUNTY_REPRESENTATIVE"
    assert mapped["sampling_site"]["source_plot_id"] == "BLAN_001"
    assert mapped["sampling_event"]["source_event_id"] == "event-a"
    assert mapped["source_geography"]["spatial_uncertainty_meters"] == 10

    for case in ("unmapped_source_only_site", "ambiguous_site_mapping"):
        record = records[case]
        assert "county_fips" not in record
        assert record["county_relationship"]["county_fips"] is None
        assert record["sampling_site"]["source_site_id"]
        assert record["sampling_event"]["source_event_id"]

    repeated = records["repeated_event_same_site"]
    replicate = records["multiple_replicates"]
    assert repeated["sampling_site"] == mapped["sampling_site"]
    assert (
        repeated["sampling_event"]["source_event_id"] != mapped["sampling_event"]["source_event_id"]
    )
    assert replicate["sampling_event"]["source_replicate_id"] == "replicate-2"


def test_site_event_schema_rejects_county_fabrication_and_missing_lineage() -> None:
    schema = json.loads(
        Path(
            "docs/contracts/tick-surveillance/canonical-tick-surveillance-v1.schema.json"
        ).read_text(encoding="utf-8")
    )
    record = json.loads(
        Path("tests/fixtures/tick_surveillance/site-event-observations.json").read_text(
            encoding="utf-8"
        )
    )[1]["record"]
    validator = Draft202012Validator(schema)

    missing_event = dict(record)
    missing_event.pop("sampling_event")
    assert list(validator.iter_errors(missing_event))

    county_status = json.loads(
        Path("tests/fixtures/tick_surveillance/site-event-observations.json").read_text(
            encoding="utf-8"
        )
    )[0]["record"]
    county_status.pop("county_fips")
    assert list(validator.iter_errors(county_status))

    fabricated_county = json.loads(json.dumps(record))
    fabricated_county["county_relationship"]["mapping_status"] = "UNMAPPED"
    fabricated_county["county_relationship"]["county_fips"] = "51061"
    assert list(validator.iter_errors(fabricated_county))


def test_canonical_site_event_identity_is_stable_and_does_not_collapse_replicates() -> None:
    common = {
        "source_dataset_id": "DP1.10093.001",
        "data_source_version_id": "RELEASE-2026",
        "source_record_id": "release-2026:row-1",
        "observation_type": "COLLECTION_ABUNDANCE",
        "sampling_site_id": "BLAN:BLAN_001",
        "sampling_event_id": "event-a",
        "sample_id": "sample-a",
        "subsample_id": "sub-a",
        "strata": {"tick_species": "source-native-taxon", "life_stage": "source-native-stage"},
    }
    first = canonical_observation_id(**common, replicate_id="replicate-1")
    assert first == canonical_observation_id(**common, replicate_id="replicate-1")
    assert first != canonical_observation_id(**common, replicate_id="replicate-2")
    different_event = {**common, "sampling_event_id": "event-b"}
    assert first != canonical_observation_id(**different_event, replicate_id="replicate-1")
    different_strata = {
        **common,
        "strata": {
            "tick_species": "another-source-native-taxon",
            "life_stage": "source-native-stage",
        },
    }
    assert first != canonical_observation_id(**different_strata, replicate_id="replicate-1")


def test_tick_profile_is_governed_envelope_only_and_pins_first_party_workbook() -> None:
    source = profile()
    assert source["resource_key"] == tick.RESOURCE_KEY
    assert source["connector_name"] == "HTTP_XLSX_V1"
    assert source["onboarding_environments"] == ["dev", "prod"]
    assert str(source["endpoint_template"]).startswith("https://www.cdc.gov/")
    assert source["deterministic_order_clause"] == "FIPSCode ASC"


def test_workbook_parser_preserves_status_and_bounded_fips_order() -> None:
    evidence = tick._parse_workbook(workbook_bytes(), profile(), 25)
    assert evidence.row_count == 30
    assert len(evidence.sample) == 25
    assert evidence.sample[0]["FIPSCode"] == "01001"
    assert evidence.sample[1]["Ixodes_scapularis_County_Status"] == "No records"
    assert evidence.schema["full_dataset_quality_validated"] is False
    assert evidence.schema["headers"] == list(tick.EXPECTED_HEADERS)
    assert evidence.schema["embedded_data_use_agreement_validated"] is True
    assert evidence.schema["embedded_classification_terms_validated"] is True


def test_workbook_parser_validates_rows_beyond_serialized_sample() -> None:
    rows = [
        (
            f"{1000 + index:05d}",
            "Fixture State",
            f"Fixture County {index}",
            "Reported",
            "Fixture citation",
            "Reported" if index < 30 else "Absent",
            "Fixture citation",
        )
        for index in range(1, 31)
    ]

    with pytest.raises(ValueError, match="unreviewed county-status"):
        tick._parse_workbook(workbook_bytes(rows=rows), profile(), 25)


def test_restricted_tick_rows_preserve_source_statuses_and_row_hashes() -> None:
    evidence, rows = tick.restricted_tick_rows(workbook_bytes(), profile())

    assert evidence.row_count == 30
    assert len(rows) == 30
    assert rows[0].fips == "01001"
    assert rows[0].scapularis_status == "Established"
    assert rows[0].pacificus_status == "Reported"
    assert len(rows[0].source_row_hash) == 64
    assert rows[0].raw["FIPSCode"] == "01001"


def test_restricted_tick_derivation_requires_protected_prod_envelope(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_operator_evidence_bundle(tmp_path, monkeypatch)
    monkeypatch.setattr(
        tick,
        "PipelineSettings",
        lambda: SimpleNamespace(topx_env="prod", enable_production_execution=False),
    )
    connection = MagicMock()
    monkeypatch.setattr(tick, "connect", lambda _: connection)

    with pytest.raises(ValueError, match="protected operator envelope"):
        tick.ingest_restricted_tick(
            evidence_bundle_dir=bundle_dir,
            evidence_run_id="12345678-1234-4234-8234-123456789abc",
        )

    connection.assert_not_called()


def test_restricted_tick_derivation_uses_owner_rights_procedure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_operator_evidence_bundle(tmp_path, monkeypatch)
    monkeypatch.setenv("RESTRICTED_CDC_PROD_OPERATOR_ENVELOPE", "true")
    monkeypatch.setattr(
        tick,
        "PipelineSettings",
        lambda: SimpleNamespace(topx_env="prod", enable_production_execution=True),
    )
    connection = MagicMock()
    connection.__enter__.return_value = connection
    cursor = connection.cursor.return_value.__enter__.return_value
    cursor.fetchone.return_value = ({"semantic_release_state": "PENDING_FINAL_COPY"},)
    monkeypatch.setattr(tick, "connect", lambda _: connection)

    result = tick.ingest_restricted_tick(
        evidence_bundle_dir=bundle_dir,
        evidence_run_id="12345678-1234-4234-8234-123456789abc",
    )

    assert result["status"] == "STAGED"
    statements = [call.args[0] for call in cursor.execute.call_args_list]
    assert "RESTRICTED_FULL_PROD" in statements[0]
    assert "CALL GOVERNANCE.SP_LOAD_RESTRICTED_TICK_PROD" in statements[1]
    assert "RAW.RESTRICTED" not in "\n".join(statements)
    connection.commit.assert_called_once()


@pytest.mark.parametrize(
    ("headers", "rows", "message"),
    [
        (("Changed", *tick.EXPECTED_HEADERS[1:]), None, "schema changed"),
        (
            tick.EXPECTED_HEADERS,
            [
                ("01002", "State", "County 2", "Reported", "Citation", "Reported", "Citation"),
                ("01001", "State", "County 1", "Reported", "Citation", "Reported", "Citation"),
            ],
            "not deterministically ordered",
        ),
        (
            tick.EXPECTED_HEADERS,
            [("1001", "State", "County", "Reported", "Citation", "Reported", "Citation")],
            "five-character county FIPS",
        ),
        (
            tick.EXPECTED_HEADERS,
            [("01001", "State", "County", "Absent", "Citation", "Reported", "Citation")],
            "unreviewed county-status",
        ),
    ],
)
def test_workbook_parser_fails_closed_on_unreviewed_structure(
    headers: tuple[str, ...], rows: list[tuple[object, ...]] | None, message: str
) -> None:
    with pytest.raises(ValueError, match=message):
        tick._parse_workbook(
            workbook_bytes(headers=headers, rows=rows), profile(), min(2, len(rows or [1, 2]))
        )


def test_tick_evidence_rejects_unprotected_prod_before_network_or_storage(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(tick, "PipelineSettings", lambda: SimpleNamespace(topx_env="prod"))
    fetch = MagicMock()
    monkeypatch.setattr(tick, "_fetch_bytes", fetch)
    with pytest.raises(ValueError, match="protected operator envelope"):
        tick.collect_tick_surveillance_evidence()
    fetch.assert_not_called()


def test_cdc_fetch_uses_browser_compatible_first_party_request_headers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    response = MagicMock()
    response.__enter__.return_value = response
    response.status_code = 200
    response.url = "https://www.cdc.gov/ticks/example.html"
    response.headers = {"content-type": "text/html", "content-length": "2"}
    response.iter_bytes.return_value = iter([b"ok"])
    client = MagicMock()
    client.__enter__.return_value = client
    client.stream.return_value = response
    monkeypatch.setattr(tick.httpx, "Client", lambda **_kwargs: client)

    result = tick._fetch_bytes(
        "https://www.cdc.gov/ticks/example.html",
        accept="text/html",
        maximum_bytes=100,
    )

    assert result.payload == b"ok"
    headers = client.stream.call_args.kwargs["headers"]
    assert headers["User-Agent"] == tick.BROWSER_USER_AGENT
    assert headers["Accept-Language"] == "en-US,en;q=0.9"
    assert headers["Cache-Control"] == "no-cache"


def test_tick_evidence_retains_non_retryable_http_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = SimpleNamespace(topx_env="dev")
    monkeypatch.setattr(tick, "PipelineSettings", lambda: settings)
    request = tick.httpx.Request("GET", str(profile()["landing_page_url"]))
    response = tick.httpx.Response(403, request=request)
    monkeypatch.setattr(
        tick,
        "_fetch_bytes",
        MagicMock(
            side_effect=tick.httpx.HTTPStatusError("forbidden", request=request, response=response)
        ),
    )
    connection = MagicMock()
    connection.__enter__.return_value = connection
    monkeypatch.setattr(tick, "connect", lambda _: connection)

    with pytest.raises(tick.httpx.HTTPStatusError):
        tick.collect_tick_surveillance_evidence()

    cursor = connection.cursor.return_value.__enter__.return_value
    statements = [call.args[0] for call in cursor.execute.call_args_list]
    assert "'RUNNING'" in statements[0]
    assert "status='FAILED'" in statements[-1]
    assert cursor.execute.call_args_list[-1].args[1][1] == "HTTP_403"
    assert "forbidden" not in statements[-1]
    connection.rollback.assert_called_once()
    assert connection.commit.call_count == 2


def test_evidence_bundle_verifies_checksums_and_runtime_provenance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle = tick._load_evidence_bundle(write_evidence_bundle(tmp_path, monkeypatch), profile())

    assert bundle.landing.media_type == "text/html"
    assert bundle.workbook.media_type == tick.XLSX_MEDIA_TYPE
    assert bundle.manifest["acquisition_route"] == tick.EVIDENCE_ROUTE


def test_operator_evidence_bundle_preserves_unobserved_http_status(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle = tick._load_evidence_bundle(
        write_operator_evidence_bundle(tmp_path, monkeypatch), profile()
    )

    assert bundle.landing.media_type == tick.PDF_MEDIA_TYPE
    assert bundle.manifest["acquisition_route"] == tick.OPERATOR_EVIDENCE_ROUTE
    assert bundle.manifest["resources"][0]["status_code"] is None


def test_operator_evidence_bundle_fails_on_retrieval_identity_mismatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_operator_evidence_bundle(tmp_path, monkeypatch)
    monkeypatch.setenv("TICK_EVIDENCE_OPERATOR_RETRIEVAL_ID", str(tick.uuid.uuid4()))

    with pytest.raises(ValueError, match="operator provenance"):
        tick._load_evidence_bundle(bundle_dir, profile())


def test_operator_evidence_records_unknown_http_status(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = SimpleNamespace(
        topx_env="dev",
        spaces_bucket="fixture-dev",
        spaces_prefix="dev",
    )
    monkeypatch.setattr(tick, "PipelineSettings", lambda: settings)
    bundle_dir = write_operator_evidence_bundle(tmp_path, monkeypatch)
    s3 = MagicMock()
    monkeypatch.setattr(tick, "_spaces_client", lambda _: s3)
    connection = MagicMock()
    connection.__enter__.return_value = connection
    monkeypatch.setattr(tick, "connect", lambda _: connection)

    result = tick.collect_tick_surveillance_evidence(25, evidence_bundle_dir=bundle_dir)

    assert result["status"] == "PENDING_STEWARD_REVIEW"
    cursor = connection.cursor.return_value.__enter__.return_value
    request_calls = [
        call
        for call in cursor.execute.call_args_list
        if "GOVERNANCE.INGESTION_REQUESTS" in call.args[0]
    ]
    assert len(request_calls) == 2
    assert all(call.args[1][6] is None for call in request_calls)
    assert any(
        call.kwargs.get("ContentType") == tick.PDF_MEDIA_TYPE
        for call in s3.put_object.call_args_list
    )


def test_evidence_bundle_fails_closed_after_payload_tampering(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_evidence_bundle(tmp_path, monkeypatch)
    (bundle_dir / "workbook.xlsx").write_bytes(b"tampered")

    with pytest.raises(ValueError, match="checksum or byte count"):
        tick._load_evidence_bundle(bundle_dir, profile())


def test_evidence_bundle_fails_closed_on_runtime_digest_mismatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_evidence_bundle(tmp_path, monkeypatch)
    monkeypatch.setenv("TICK_EVIDENCE_BASE_IMAGE_DIGEST", "sha256:" + "d" * 64)

    with pytest.raises(ValueError, match="base image digest"):
        tick._load_evidence_bundle(bundle_dir, profile())


def test_capture_failure_details_are_stage_aware_and_redacted() -> None:
    classification, diagnostic = tick._failure_details(
        ValueError("SPACES_SECRET_ACCESS_KEY=not-for-the-ledger"),
        "private_artifact_retention",
    )

    assert classification == "EVIDENCE_CAPTURE_FAILED"
    assert diagnostic == (
        "CDC tick-surveillance evidence capture failed during "
        "private_artifact_retention; error_type=ValueError; review protected logs"
    )
    assert "SPACES_SECRET_ACCESS_KEY" not in diagnostic
    assert "not-for-the-ledger" not in diagnostic

    validation_classification, _ = tick._failure_details(
        ValueError("schema changed"), "source_validation"
    )
    assert validation_classification == "SOURCE_VALIDATION_FAILED"


def test_evidence_bundle_rejects_duplicate_manifest_keys(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bundle_dir = write_evidence_bundle(tmp_path, monkeypatch)
    payload = (bundle_dir / "acquisition-manifest.json").read_text()
    (bundle_dir / "acquisition-manifest.json").write_text(
        payload.replace('{"manifest_version": 1,', '{"manifest_version": 1, "manifest_version": 1,')
    )

    with pytest.raises(ValueError, match="duplicate key"):
        tick._load_evidence_bundle(bundle_dir, profile())


def test_tick_evidence_creates_only_pending_review_evidence(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = SimpleNamespace(
        topx_env="dev",
        spaces_bucket="fixture-dev",
        spaces_prefix="dev",
    )
    monkeypatch.setattr(tick, "PipelineSettings", lambda: settings)
    bundle_dir = write_evidence_bundle(tmp_path, monkeypatch)
    fetch = MagicMock()
    monkeypatch.setattr(tick, "_fetch_bytes", fetch)
    monkeypatch.setattr(tick, "_spaces_client", lambda _: object())

    def save_artifact(
        _s3: object,
        local_settings: object,
        run_id: str,
        payload: bytes,
        _media_type: str,
        resource_key: str = tick.RESOURCE_KEY,
    ) -> object:
        return create_artifact(
            payload=payload,
            environment=str(local_settings.topx_env),
            resource_key=resource_key,
            run_id=run_id,
        )

    monkeypatch.setattr(tick, "_save_artifact", save_artifact)
    connection = MagicMock()
    connection.__enter__.return_value = connection
    monkeypatch.setattr(tick, "connect", lambda _: connection)

    result = tick.collect_tick_surveillance_evidence(25, evidence_bundle_dir=bundle_dir)

    assert result["status"] == "PENDING_STEWARD_REVIEW"
    assert result["sample_rows"] == 25
    assert result["workbook_rows"] == 30
    assert result["full_dataset_quality_validated"] is False
    assert result["acquisition_manifest_sha256"] is not None
    fetch.assert_not_called()
    statements = [
        call.args[0]
        for call in connection.cursor.return_value.__enter__.return_value.execute.call_args_list
    ]
    rendered = "\n".join(statements)
    for required in (
        "GOVERNANCE.INGESTION_RUNS",
        "GOVERNANCE.CATALOG_DATASETS",
        "GOVERNANCE.CATALOG_RESOURCES",
        "GOVERNANCE.SOURCE_ACCESS_PROFILES",
        "GOVERNANCE.INGESTION_REQUESTS",
        "GOVERNANCE.RAW_ARTIFACTS",
        "GOVERNANCE.SOURCE_DOCUMENT_SNAPSHOTS",
        "GOVERNANCE.SCHEMA_SNAPSHOTS",
        "GOVERNANCE.DATASET_QUALITY_ASSESSMENTS",
    ):
        assert required in rendered
    assert any(
        "ACQUISITION_MANIFEST" in call.args[1]
        for call in connection.cursor.return_value.__enter__.return_value.execute.call_args_list
        if len(call.args) > 1 and isinstance(call.args[1], tuple)
    )
    for forbidden in (
        "DATA_SOURCE_VERSIONS",
        "MANUAL_REVIEW_DECISIONS",
        "RAW.CDC_TICK",
        "STAGING.",
        "CONFORMED.",
        "ANALYTICS.",
        "FEATURE_STORE.",
        "COPY INTO",
    ):
        assert forbidden not in rendered
    assert connection.commit.call_count == 2
    connection.rollback.assert_not_called()


def test_dev_tick_review_migrations_and_workflow_preserve_scope() -> None:
    migrations = {migration.version: migration for migration in load_migrations()}
    assert migration_execution_role(migrations["V053"], DEV_DATABASE) == (
        "OH_LYME_DEV_STREAMLIT_OWNER"
    )
    assert "cdc_tick_ixodes_county_status" in migrations["V053"].source
    assert "No records is not evidence of absence" in migrations["V053"].source
    assert "HTTP_XLSX_V1" in migrations["V054"].source
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migrations["V053"], "ONE_HEALTH_LYME_GAP_ATLAS_PROD")
    assert "V053" not in {
        item["version"] for item in migration_plan("ONE_HEALTH_LYME_GAP_ATLAS_PROD")
    }
    assert not Path(".github/workflows/capture-dev-cdc-tick-surveillance.yml").exists()

    operator_workflow = Path(
        ".github/workflows/capture-dev-cdc-tick-surveillance-operator.yml"
    ).read_text(encoding="utf-8")
    assert "TICK_EVIDENCE_OPERATOR_RETRIEVAL_ID" in operator_workflow
    assert "delete-tag pipeline" in operator_workflow
    assert "cdc-tick-surveillance-sample --sample-limit 25" in operator_workflow
    assert "--evidence-bundle-dir /run/atlas-tick-evidence" in operator_workflow
    assert "TICK_EVIDENCE_BASE_IMAGE_DIGEST" in operator_workflow
    assert "TICK_EVIDENCE_ENVELOPE_DIGEST" in operator_workflow
    assert "PRE_DEPLOY" in operator_workflow
    assert ".jobs |= map(.image.digest = $envelope_digest)" in operator_workflow
    assert "[.jobs[].image.digest] | unique | length" in operator_workflow
    assert 'sort -u)" = "$ENVELOPE_DIGEST"' in operator_workflow
    assert "SNOWFLAKE_" not in operator_workflow
    assert "SPACES_" not in operator_workflow
    for forbidden in ("ingest-approved", "dbt", "PROD_APP_ID", "production"):
        assert forbidden not in operator_workflow


def test_approval_console_exposes_tick_candidate_only_in_dev() -> None:
    source = Path("streamlit_approval/streamlit_app.py").read_text(encoding="utf-8")
    assert 'if current_database.endswith("_DEV")' in source
    assert 'source_labels["cdc_tick_ixodes_county_status"]' in source
    assert "No records means no reported surveillance evidence" in source
    assert "implemented and authorized acquisition path" in source
    assert "INSERT INTO GOVERNANCE" not in source
    assert "UPDATE GOVERNANCE" not in source
