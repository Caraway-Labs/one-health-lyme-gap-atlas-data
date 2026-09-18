from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from lyme_gap_atlas_data import pathogen_surveillance as pathogen


def workbook_bytes(
    *,
    headers: tuple[str, ...] = pathogen.EXPECTED_HEADERS,
    rows: list[tuple[object, ...]] | None = None,
) -> bytes:
    workbook = Workbook()
    agreement = workbook.active
    agreement.title = "Data Use Agreement"
    agreement.append(["Access to ArboNET Tick Module data is limited to the Requestor."])
    agreement.append(["These data should not be provided to other persons."])
    agreement.append(["ArboNET will be appropriately referenced."])
    agreement.append(["A final copy of publications will be provided to CDC."])
    terms = workbook.create_sheet("Classification Terms")
    terms.append(["County classification", "Definition"])
    terms.append(["Present", "A pathogen was identified in host-seeking Ixodes ticks."])
    terms.append(["No records", "This should not be interpreted as pathogen absence."])
    data = workbook.create_sheet("Ixodes Pathogens 2025")
    data.append(["Tickborne pathogens through Dec. 31, 2025"])
    data.append(list(headers))
    source_rows = rows or [
        (
            f"{1000 + index:05d}",
            "Fixture State",
            f"Fixture County {index}",
            *sum(
                (
                    ["Present" if index % 2 else "No records", f"Fixture source {column}"]
                    for column in range(7)
                ),
                [],
            ),
        )
        for index in range(1, 31)
    ]
    for row in source_rows:
        data.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_pathogen_profile_is_private_governed_envelope_evidence_only() -> None:
    profile = pathogen.load_pathogen_profile()
    assert profile["resource_key"] == pathogen.RESOURCE_KEY
    assert profile["onboarding_environments"] == ["dev", "prod"]
    assert profile["onboarding_mode"] == "EVIDENCE_ONLY"
    assert profile["allowed_status_values"] == ["Present", "No records"]
    assert profile["landing_page_url"].startswith("https://www.cdc.gov/")


def test_parser_preserves_distinct_pathogen_status_semantics() -> None:
    evidence = pathogen.parse_pathogen_workbook(
        workbook_bytes(), pathogen.load_pathogen_profile(), sample_limit=25
    )
    assert evidence.valid_county_row_count == 30
    assert evidence.blank_fips_row_count == 0
    assert len(evidence.sample) == 25
    assert evidence.sample[0]["FIPS_Code"] == "01001"
    assert evidence.schema["observation_type"] == "PATHOGEN_PRESENCE_STATUS"
    assert evidence.schema["embedded_data_use_agreement_validated"] is True
    assert evidence.schema["full_dataset_quality_validated"] is False
    assert evidence.schema["dataset_as_of"] == "2025-12-31"
    json.dumps(evidence.schema, sort_keys=True, separators=(",", ":"))


def test_parser_counts_blank_fips_rows_but_rejects_invalid_values() -> None:
    row = ("01001", "State", "County", *sum((["Present", "source"] for _ in range(7)), []))
    blank = (None, "", "", *sum((["No records", "source"] for _ in range(7)), []))
    evidence = pathogen.parse_pathogen_workbook(
        workbook_bytes(rows=[row, blank]), pathogen.load_pathogen_profile(), sample_limit=1
    )
    assert evidence.valid_county_row_count == 1
    assert evidence.blank_fips_row_count == 1
    malformed = (
        "not-a-fips",
        "State",
        "County",
        *sum((["Present", "source"] for _ in range(7)), []),
    )
    with pytest.raises(ValueError, match="five-character county FIPS"):
        pathogen.parse_pathogen_workbook(
            workbook_bytes(rows=[malformed]), pathogen.load_pathogen_profile(), sample_limit=1
        )


def test_restricted_rows_keep_source_faithful_lineage_inside_private_loader() -> None:
    row = (
        "01001",
        "State",
        "County",
        *sum((["Present", "Restricted source category"] for _ in range(7)), []),
    )
    evidence, rows = pathogen.restricted_pathogen_rows(
        workbook_bytes(rows=[row]), pathogen.load_pathogen_profile()
    )
    assert evidence.valid_county_row_count == 1
    assert len(rows) == 1
    assert rows[0].fips == "01001"
    assert rows[0].burgdorferi_status == "Present"
    assert rows[0].burgdorferi_source == "Restricted source category"
    assert len(rows[0].source_row_hash) == 64
    assert rows[0].raw["State"] == "State"


def test_parser_rejects_status_relabeling_and_schema_changes() -> None:
    bad_status = ("01001", "State", "County", *sum((["Absent", "source"] for _ in range(7)), []))
    with pytest.raises(ValueError, match="unreviewed pathogen-status"):
        pathogen.parse_pathogen_workbook(
            workbook_bytes(rows=[bad_status]), pathogen.load_pathogen_profile(), sample_limit=1
        )
    changed_headers = ("Changed", *pathogen.EXPECTED_HEADERS[1:])
    with pytest.raises(ValueError, match="schema changed"):
        pathogen.parse_pathogen_workbook(
            workbook_bytes(headers=changed_headers),
            pathogen.load_pathogen_profile(),
            sample_limit=1,
        )


def test_pathogen_capture_contract_is_private_and_profile_bound() -> None:
    workflow = Path(".github/workflows/capture-dev-cdc-tick-surveillance-operator.yml").read_text(
        encoding="utf-8"
    )
    cli = Path("src/lyme_gap_atlas_data/cli.py").read_text(encoding="utf-8")
    publisher = Path("scripts/publish_tick_operator_evidence.py").read_text(encoding="utf-8")
    assert "source_kind" in workflow
    assert "cdc-pathogen-surveillance-sample" in workflow
    assert '"$SOURCE_KIND-operator-evidence-$RETRIEVAL_ID"' in workflow
    assert "cdc-pathogen-surveillance-sample" in cli
    assert "--source-kind" in publisher
    assert "--operation" in publisher
    assert '"-f",\n                f"operation={args.operation}"' in publisher
    assert 'args.operation == "derive" and args.source_kind != "pathogen"' not in publisher
    assert "private restricted derivation" in publisher
    assert "GitHub artifact" not in workflow


def test_pathogen_review_console_is_dev_only_and_does_not_expose_workbook() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_execution_role,
        migration_plan,
        render_migration,
    )

    migrations = {migration.version: migration for migration in load_migrations()}
    review = migrations["V074"]
    assert migration_execution_role(review, DEV_DATABASE) is None
    assert "cdc_tick_ixodes_pathogen_status" in review.source
    assert "requestor-restricted" in review.source
    assert "RAW_ARTIFACTS" not in review.source
    assert "No records is not pathogen absence" in review.source
    procedure = migrations["V075"]
    assert migration_execution_role(procedure, DEV_DATABASE) is None
    assert "cdc_tick_ixodes_pathogen_status" in procedure.source
    assert (
        "GRANT USAGE ON PROCEDURE GOVERNANCE.SP_RECORD_SOURCE_REVIEW_DECISION" in procedure.source
    )
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(review, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")
    assert "V074" not in {
        item["version"] for item in migration_plan("ONE_HEALTH_LYME_GAP_ATLAS_PROD")
    }
    assert "V075" not in {
        item["version"] for item in migration_plan("ONE_HEALTH_LYME_GAP_ATLAS_PROD")
    }

    console = Path("streamlit_approval/streamlit_app.py").read_text(encoding="utf-8")
    assert 'source_labels["cdc_tick_ixodes_pathogen_status"]' in console
    assert "private workbook is not displayed, exported, or published here" in console


def test_restricted_derivation_boundary_is_dev_only_and_procedure_only_for_runtime() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V077"]
    source = migration.source
    assert "RAW.RESTRICTED_CDC_PATHOGEN_WORKBOOK_ROWS" in source
    assert "STAGING.RESTRICTED_CDC_PATHOGEN_WORKBOOK_ROWS" in source
    assert "CONFORMED.RESTRICTED_CDC_PATHOGEN_COUNTY_STATUS" in source
    assert "EXECUTE AS OWNER" in source
    assert "GRANT USAGE ON PROCEDURE" in source
    assert "GRANT SELECT ON TABLE RAW.RESTRICTED" not in source
    assert "GRANT SELECT ON TABLE STAGING.RESTRICTED" not in source
    assert "GRANT SELECT ON TABLE CONFORMED.RESTRICTED" not in source
    assert "V077" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_restricted_derivation_owner_can_record_its_terminal_run_state_only_in_dev() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V078"]
    assert (
        "GRANT UPDATE ON TABLE GOVERNANCE.INGESTION_RUNS TO ROLE "
        "OH_LYME_DEV_MIGRATION_DEPLOYER" in migration.source
    )
    assert "V078" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_restricted_derivation_owner_can_write_only_required_governance_outputs_in_dev() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V079"]
    for table_name in (
        "INGESTION_REQUESTS",
        "DATA_QUALITY_RESULTS",
        "INGESTION_PUBLICATIONS",
    ):
        assert (
            f"GRANT INSERT ON TABLE GOVERNANCE.{table_name} TO ROLE "
            "OH_LYME_DEV_MIGRATION_DEPLOYER" in migration.source
        )
    assert "V079" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_restricted_pathogen_parity_classification_is_dev_only_and_unknown() -> None:
    from lyme_gap_atlas_data.migrations import DEV_DATABASE, load_migrations, migration_plan

    migration = {item.version: item for item in load_migrations()}["V080"]
    assert "UNKNOWN_SOURCE_COVERAGE" in migration.source
    assert "SP_CLASSIFY_RESTRICTED_PATHOGEN_PARITY_DEV" in migration.source
    assert "V080" in {item["version"] for item in migration_plan(DEV_DATABASE)}


def test_evidence_only_tick_coverage_classification_is_dev_only_and_unknown() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V081"]
    assert "UNKNOWN_SOURCE_COVERAGE" in migration.source
    assert "SP_CLASSIFY_EVIDENCE_ONLY_TICK_COVERAGE_DEV" in migration.source
    assert "V081" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_evidence_only_tick_coverage_owner_reads_only_its_validation_ledgers() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V082"]
    for table_name in ("INGESTION_RUNS", "DATA_SOURCE_VERSIONS"):
        assert (
            f"GRANT SELECT ON TABLE GOVERNANCE.{table_name} TO ROLE "
            "OH_LYME_DEV_MIGRATION_DEPLOYER" in migration.source
        )
    assert "RESTRICTED_CDC_PATHOGEN" not in migration.source
    assert "V082" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_semantic_runtime_reads_only_the_tick_coverage_classification() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V083"]
    assert (
        "GRANT SELECT ON TABLE GOVERNANCE.EVIDENCE_ONLY_SOURCE_COVERAGE_CLASSIFICATIONS"
        in migration.source
    )
    assert "OH_LYME_DEV_RUNTIME" in migration.source
    assert "RESTRICTED_CDC_PATHOGEN" not in migration.source
    assert "GOVERNED_SOURCE_RECORDS" not in migration.source
    assert "V083" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_semantic_runtime_reads_only_the_pathogen_parity_classification() -> None:
    from lyme_gap_atlas_data.migrations import (
        DEV_DATABASE,
        load_migrations,
        migration_plan,
        render_migration,
    )

    migration = {item.version: item for item in load_migrations()}["V084"]
    assert (
        "GRANT SELECT ON TABLE GOVERNANCE.RESTRICTED_PATHOGEN_PARITY_CLASSIFICATIONS"
        in migration.source
    )
    assert "OH_LYME_DEV_RUNTIME" in migration.source
    assert "RESTRICTED_CDC_PATHOGEN_COUNTY_STATUS" not in migration.source
    assert "GOVERNED_SOURCE_RECORDS" not in migration.source
    assert "V084" in {item["version"] for item in migration_plan(DEV_DATABASE)}
    with pytest.raises(ValueError, match="DEV-only"):
        render_migration(migration, "ONE_HEALTH_LYME_GAP_ATLAS_PROD")


def test_pathogen_derivation_workflow_requires_explicit_private_operation() -> None:
    workflow = Path(".github/workflows/capture-dev-cdc-tick-surveillance-operator.yml").read_text(
        encoding="utf-8"
    )
    assert "operation:" in workflow
    assert "options: [evidence, derive]" in workflow
    assert '"$OPERATION" != "derive" || "$SOURCE_KIND" = "pathogen"' in workflow
    assert "cdc-pathogen-restricted-dev-capture-and-ingest" in workflow


def test_production_restricted_operator_path_is_protected_and_restores_topology() -> None:
    from lyme_gap_atlas_data.migrations import (
        PROD_DATABASE,
        load_migrations,
        migration_execution_role,
        migration_plan,
    )

    workflow = Path(".github/workflows/capture-prod-cdc-restricted-operator.yml").read_text(
        encoding="utf-8"
    )
    migration = {item.version: item for item in load_migrations()}["V087"]
    tick_migration = {item.version: item for item in load_migrations()}["V088"]
    release_gate_migration = {item.version: item for item in load_migrations()}["V089"]
    caller_rights_migration = {item.version: item for item in load_migrations()}["V090"]
    variant_binding_migration = {item.version: item for item in load_migrations()}["V091"]
    select_binding_migration = {item.version: item for item in load_migrations()}["V092"]
    assert "environment: production" in workflow
    assert "PROD_APP_ID: ${{ vars.PROD_APP_ID }}" in workflow
    assert "RESTRICTED_CDC_PROD_OPERATOR_ENVELOPE" in workflow
    assert "delete-tag pipeline" in workflow
    assert "approved-source-ingestion" in workflow
    assert "GitHub artifact" not in workflow
    assert "SP_LOAD_RESTRICTED_PATHOGEN_PROD" in migration.source
    assert "SP_LOAD_RESTRICTED_TICK_PROD" in tick_migration.source
    assert "RESTRICTED_CDC_TICK_COUNTY_STATUS" in tick_migration.source
    assert "SP_RECORD_RESTRICTED_FINAL_COPY_PROD" in release_gate_migration.source
    assert "RESTRICTED_PATHOGEN_PARITY_CLASSIFICATIONS" in release_gate_migration.source
    assert "SP_RECORD_RESTRICTED_SOURCE_REVIEW_PROD" in caller_rights_migration.source
    assert "EXECUTE AS CALLER" in caller_rights_migration.source
    assert (
        "GRANT UPDATE ON TABLE GOVERNANCE.DATA_SOURCE_VERSIONS"
        not in caller_rights_migration.source
    )
    assert "EXECUTE AS CALLER" in variant_binding_migration.source
    assert (
        "SELECT :decision_id,:RESOURCE_KEY,:DECISION,:RATIONALE,CONDITIONS"
        in variant_binding_migration.source
    )
    assert (
        "VALUES(:decision_id,:RESOURCE_KEY,:DECISION,:RATIONALE,:CONDITIONS"
        not in variant_binding_migration.source
    )
    assert (
        "SELECT :decision_id,:RESOURCE_KEY,:DECISION,:RATIONALE,:CONDITIONS"
        in select_binding_migration.source
    )
    assert "GRANT SELECT ON TABLE RAW.RESTRICTED" not in migration.source
    assert "RESTRICTED_SOURCE_PUBLICATION_ATTESTATIONS" in migration.source
    assert "V087" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "V088" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "V089" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "V090" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "V091" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "V092" in {item["version"] for item in migration_plan(PROD_DATABASE)}
    assert "cdc-tick-restricted-prod-ingest" in workflow
    assert '"SNOWFLAKE_ROLE",scope:"RUN_TIME",value:"OH_LYME_PROD_RUNTIME"' in workflow
    assert '"$OPERATION" != derive || "$SOURCE_KIND" = pathogen' not in workflow
    assert migration_execution_role(migration, PROD_DATABASE) is None


def test_restricted_operator_publisher_passes_the_derivation_evidence_input() -> None:
    publisher = Path("scripts/publish_tick_operator_evidence.py").read_text(encoding="utf-8")

    assert '["-f", f"evidence_run_id={args.evidence_run_id}"]' in publisher
