"""Adapter protocol and registry for simplified ingestion."""

from __future__ import annotations

import hashlib
import io
import json
import os
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

import httpx
from openpyxl import load_workbook  # type: ignore[import-untyped]

from .types import AdapterKind, FailureCategory, SourceDefinition, ValidationIssue, ValidationResult


@dataclass
class AcquireResult:
    payload: Any
    artifact_sha256: str
    media_type: str
    row_count: int | None = None
    detail: dict[str, Any] | None = None
    raw_payload: bytes | None = None


class AcquisitionError(RuntimeError):
    """A bounded provider/network failure safe to expose as a diagnostic code."""

    def __init__(self, message: str, *, code: str) -> None:
        self.code = code
        super().__init__(message)


@dataclass
class NormalizeResult:
    records: list[dict[str, Any]]
    transformation_version: str
    detail: dict[str, Any] | None = None


class SourceAdapter(Protocol):
    kind: AdapterKind

    def acquire(
        self,
        definition: SourceDefinition,
        *,
        fixture_dir: Path | None = None,
    ) -> AcquireResult: ...

    def validate_payload(self, definition: SourceDefinition, payload: Any) -> ValidationResult: ...

    def normalize(self, definition: SourceDefinition, payload: Any) -> NormalizeResult: ...

    def restore_raw_payload(self, definition: SourceDefinition, raw_payload: bytes) -> Any: ...


class SocrataAdapter:
    """Socrata SODA2 adapter for the x5j9 golden path."""

    kind = AdapterKind.SOCRATA

    def __init__(
        self,
        *,
        client: httpx.Client | None = None,
        sleep_fn: Any = time.sleep,
        max_retries: int = 3,
    ) -> None:
        self._client = client
        self._sleep = sleep_fn
        self._max_retries = max_retries

    def acquire(
        self,
        definition: SourceDefinition,
        *,
        fixture_dir: Path | None = None,
    ) -> AcquireResult:
        if fixture_dir is not None:
            sample_path = fixture_dir / "sample.json"
            metadata_path = fixture_dir / "metadata.json"
            sample = json.loads(sample_path.read_text(encoding="utf-8"))
            fixture_metadata = (
                json.loads(metadata_path.read_text(encoding="utf-8"))
                if metadata_path.exists()
                else {}
            )
            payload = {"metadata": fixture_metadata, "sample": sample}
            body = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str).encode(
                "utf-8"
            )
            return AcquireResult(
                payload=payload,
                artifact_sha256=hashlib.sha256(body).hexdigest(),
                media_type="application/json",
                row_count=len(sample) if isinstance(sample, list) else None,
                detail={"source": "fixture"},
                raw_payload=body,
            )
        headers = {"Accept": "application/json"}
        if definition.auth_mode.value == "app_token_env":
            env_name = str(definition.extra.get("auth_env_var") or "SOCRATA_APP_TOKEN")
            token = os.environ.get(env_name)
            if not token:
                raise AcquisitionError(
                    f"{env_name} is required for this SourceDefinition", code="AUTH_REQUIRED"
                )
            headers["X-App-Token"] = token
        client = self._client or httpx.Client(timeout=30.0, follow_redirects=False)
        try:
            metadata_payload: Any = {}
            raw_parts: list[bytes] = []
            if definition.metadata_endpoint_template:
                metadata_response = self._request(
                    client, definition.metadata_endpoint_template, headers
                )
                try:
                    metadata_payload = metadata_response.json()
                except ValueError as error:
                    raise AcquisitionError(
                        "Socrata metadata was not valid JSON", code="RESPONSE_JSON"
                    ) from error
                raw_parts.append(metadata_response.content)

            offset = 0
            rows: list[dict[str, Any]] = []
            pages: list[dict[str, Any]] = []
            while True:
                if definition.maximum_rows is not None and len(rows) >= definition.maximum_rows:
                    break
                limit = min(
                    definition.page_size,
                    (definition.maximum_rows or definition.page_size) - len(rows),
                )
                params = {
                    "$limit": limit,
                    "$offset": offset,
                    "$order": definition.deterministic_order_clause,
                }
                response = self._request(
                    client, definition.endpoint_template, headers, params=params
                )
                try:
                    page = response.json()
                except ValueError as error:
                    raise AcquisitionError(
                        "Socrata page was not valid JSON", code="RESPONSE_JSON"
                    ) from error
                if not isinstance(page, list):
                    raise AcquisitionError(
                        "Socrata response was not a JSON array", code="RESPONSE_SHAPE"
                    )
                if not all(isinstance(row, dict) for row in page):
                    raise AcquisitionError("Socrata rows must be JSON objects", code="ROW_SHAPE")
                raw_parts.append(response.content)
                pages.append(
                    {"offset": offset, "row_count": len(page), "sha256": _sha256(response.content)}
                )
                rows.extend(page)
                if len(page) < limit:
                    break
                offset += len(page)

            return AcquireResult(
                payload={"metadata": metadata_payload, "sample": rows, "pages": pages},
                artifact_sha256=_sha256(b"\n".join(raw_parts)),
                media_type="application/octet-stream",
                row_count=len(rows),
                detail={
                    "source": "live",
                    "page_count": len(pages),
                    "request_order": definition.deterministic_order_clause,
                },
                raw_payload=b"\n".join(raw_parts),
            )
        finally:
            if self._client is None:
                client.close()

    def _request(
        self,
        client: httpx.Client,
        url: str,
        headers: dict[str, str],
        *,
        params: dict[str, Any] | None = None,
    ) -> httpx.Response:
        for attempt in range(self._max_retries):
            try:
                response = client.get(url, headers=headers, params=params)
            except httpx.TransportError as error:
                if attempt == self._max_retries - 1:
                    raise AcquisitionError("Socrata transport failed", code="TRANSPORT") from error
                self._sleep(2**attempt)
                continue
            if response.status_code in {408, 425, 429} or response.status_code >= 500:
                if attempt == self._max_retries - 1:
                    raise AcquisitionError(
                        f"Socrata provider returned HTTP {response.status_code}",
                        code=f"HTTP_{response.status_code}",
                    )
                self._sleep(2**attempt)
                continue
            if response.status_code >= 400:
                raise AcquisitionError(
                    f"Socrata provider rejected request with HTTP {response.status_code}",
                    code=f"HTTP_{response.status_code}",
                )
            return response
        raise AssertionError("unreachable")

    def validate_payload(self, definition: SourceDefinition, payload: Any) -> ValidationResult:
        issues: list[ValidationIssue] = []
        sample = payload.get("sample") if isinstance(payload, dict) else payload
        if not isinstance(sample, list) or not sample:
            issues.append(
                ValidationIssue(
                    code="EMPTY_SAMPLE", message="Socrata sample must be a non-empty list"
                )
            )
            return ValidationResult(ok=False, issues=issues)
        if not all(isinstance(row, dict) for row in sample):
            issues.append(ValidationIssue(code="ROW_SHAPE", message="Socrata rows must be objects"))
        if isinstance(sample[0], dict):
            missing = [column for column in definition.required_columns if column not in sample[0]]
            if missing:
                issues.append(
                    ValidationIssue(
                        code="MISSING_COLUMNS",
                        message=f"Missing required columns: {', '.join(missing)}",
                        category=FailureCategory.SCHEMA,
                    )
                )
            _validate_era(definition, sample, issues)
        return ValidationResult(ok=not issues, issues=issues)

    def restore_raw_payload(self, definition: SourceDefinition, raw_payload: bytes) -> Any:
        """Reconstruct the canonical adapter payload from retained JSON responses."""
        del definition
        decoder = json.JSONDecoder()
        text = raw_payload.decode("utf-8")
        documents: list[Any] = []
        offset = 0
        while offset < len(text):
            while offset < len(text) and text[offset].isspace():
                offset += 1
            if offset == len(text):
                break
            document, offset = decoder.raw_decode(text, offset)
            documents.append(document)
        if not documents:
            raise AcquisitionError("Retained Socrata artifact was empty", code="ARTIFACT_SHAPE")
        if (
            len(documents) == 1
            and isinstance(documents[0], dict)
            and isinstance(documents[0].get("sample"), list)
        ):
            return documents[0]
        metadata: Any = {}
        pages = documents
        if isinstance(documents[0], dict):
            metadata = documents[0]
            pages = documents[1:]
        rows: list[dict[str, Any]] = []
        for page in pages:
            if not isinstance(page, list) or not all(isinstance(row, dict) for row in page):
                raise AcquisitionError(
                    "Retained Socrata artifact had an invalid page", code="ARTIFACT_SHAPE"
                )
            rows.extend(page)
        return {"metadata": metadata, "sample": rows}

    def normalize(self, definition: SourceDefinition, payload: Any) -> NormalizeResult:
        sample = payload.get("sample") if isinstance(payload, dict) else payload
        assert isinstance(sample, list)
        records = [_normalized_record(definition, row) for row in sample if isinstance(row, dict)]
        return NormalizeResult(
            records=records,
            transformation_version="socrata_normalize_v1",
            detail={"record_count": len(records)},
        )


class HttpXlsxAdapter:
    """HTTP XLSX adapter for tick surveillance and similar file sources."""

    kind = AdapterKind.HTTP_XLSX

    def __init__(
        self,
        *,
        client: httpx.Client | None = None,
        sleep_fn: Any = time.sleep,
        max_retries: int = 3,
    ) -> None:
        self._client = client
        self._sleep = sleep_fn
        self._max_retries = max_retries

    def acquire(
        self,
        definition: SourceDefinition,
        *,
        fixture_dir: Path | None = None,
    ) -> AcquireResult:
        if fixture_dir is not None:
            sample_path = fixture_dir / "sample.json"
            sample = json.loads(sample_path.read_text(encoding="utf-8"))
            body = json.dumps(sample, sort_keys=True, separators=(",", ":"), default=str).encode(
                "utf-8"
            )
            return AcquireResult(
                payload={"sample": sample},
                artifact_sha256=hashlib.sha256(body).hexdigest(),
                media_type="application/json",
                row_count=len(sample) if isinstance(sample, list) else None,
                detail={"source": "fixture", "workbook_sheet": definition.workbook_sheet},
                raw_payload=body,
            )
        limit = definition.maximum_workbook_bytes
        if limit is None:
            raise AcquisitionError(
                "maximum_workbook_bytes is required for live XLSX", code="BYTE_BOUND"
            )
        headers = {"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
        client = self._client or httpx.Client(timeout=60.0, follow_redirects=False)
        try:
            response = self._request(client, definition.endpoint_template, headers)
            content_length = response.headers.get("content-length")
            if content_length is not None:
                try:
                    content_length_value = int(content_length)
                except ValueError as error:
                    raise AcquisitionError(
                        "XLSX content-length header was invalid", code="RESPONSE_HEADERS"
                    ) from error
                if content_length_value > limit:
                    raise AcquisitionError(
                        "XLSX content-length exceeds configured bound", code="BYTE_BOUND"
                    )
            body = response.content
            if len(body) > limit:
                raise AcquisitionError(
                    "XLSX response exceeds configured byte bound", code="BYTE_BOUND"
                )
            try:
                workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            except (OSError, ValueError, KeyError) as error:
                raise AcquisitionError(
                    "XLSX response was not a readable workbook", code="WORKBOOK_INVALID"
                ) from error
            try:
                if definition.workbook_sheet not in workbook.sheetnames:
                    raise AcquisitionError(
                        "Configured workbook sheet was not found", code="SHEET_NOT_FOUND"
                    )
                sheet = workbook[definition.workbook_sheet]
                header_row = definition.header_row or 1
                rows_iter = sheet.iter_rows(values_only=True)
                headers_row: tuple[Any, ...] | None = None
                for index, row in enumerate(rows_iter, start=1):
                    if index == header_row:
                        headers_row = tuple(row)
                        break
                if headers_row is None:
                    raise AcquisitionError(
                        "Configured workbook header row was not found", code="HEADER_NOT_FOUND"
                    )
                names = [str(value).strip() for value in headers_row]
                if len(names) != len(set(names)) or any(not name for name in names):
                    raise AcquisitionError(
                        "Workbook header is not unique and non-empty", code="HEADER_INVALID"
                    )
                missing = [column for column in definition.required_columns if column not in names]
                if missing:
                    raise AcquisitionError(
                        f"Workbook is missing required columns: {', '.join(missing)}",
                        code="MISSING_COLUMNS",
                    )
                rows: list[dict[str, Any]] = []
                for values in rows_iter:
                    if not any(value is not None for value in values):
                        continue
                    rows.append({name: value for name, value in zip(names, values, strict=False)})
                    if definition.maximum_rows is not None and len(rows) > definition.maximum_rows:
                        raise AcquisitionError(
                            "Workbook exceeds configured row bound", code="ROW_BOUND"
                        )
                return AcquireResult(
                    payload={
                        "sample": rows,
                        "workbook_sheet": definition.workbook_sheet,
                        "headers": names,
                    },
                    artifact_sha256=_sha256(body),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    row_count=len(rows),
                    detail={
                        "source": "live",
                        "workbook_sheet": definition.workbook_sheet,
                        "headers": names,
                    },
                    raw_payload=body,
                )
            finally:
                workbook.close()
        finally:
            if self._client is None:
                client.close()

    def _request(self, client: httpx.Client, url: str, headers: dict[str, str]) -> httpx.Response:
        for attempt in range(self._max_retries):
            try:
                response = client.get(url, headers=headers)
            except httpx.TransportError as error:
                if attempt == self._max_retries - 1:
                    raise AcquisitionError("XLSX transport failed", code="TRANSPORT") from error
                self._sleep(2**attempt)
                continue
            if response.status_code in {408, 425, 429} or response.status_code >= 500:
                if attempt == self._max_retries - 1:
                    raise AcquisitionError(
                        f"XLSX provider returned HTTP {response.status_code}",
                        code=f"HTTP_{response.status_code}",
                    )
                self._sleep(2**attempt)
                continue
            if response.status_code >= 400:
                raise AcquisitionError(
                    f"XLSX provider rejected request with HTTP {response.status_code}",
                    code=f"HTTP_{response.status_code}",
                )
            return response
        raise AssertionError("unreachable")

    def validate_payload(self, definition: SourceDefinition, payload: Any) -> ValidationResult:
        issues: list[ValidationIssue] = []
        sample = payload.get("sample") if isinstance(payload, dict) else payload
        if not isinstance(sample, list) or not sample:
            issues.append(
                ValidationIssue(
                    code="EMPTY_SAMPLE",
                    message="XLSX sample must be a non-empty list",
                    category=FailureCategory.SCHEMA,
                )
            )
            return ValidationResult(ok=False, issues=issues)
        first = sample[0]
        if not isinstance(first, dict):
            issues.append(
                ValidationIssue(
                    code="ROW_SHAPE",
                    message="XLSX rows must be objects",
                    category=FailureCategory.SCHEMA,
                )
            )
            return ValidationResult(ok=False, issues=issues)
        missing = [column for column in definition.required_columns if column not in first]
        if missing:
            issues.append(
                ValidationIssue(
                    code="MISSING_COLUMNS",
                    message=f"Missing required columns: {', '.join(missing)}",
                    category=FailureCategory.SCHEMA,
                )
            )
        if isinstance(sample[0], dict):
            _validate_era(definition, sample, issues)
        return ValidationResult(ok=not issues, issues=issues)

    def restore_raw_payload(self, definition: SourceDefinition, raw_payload: bytes) -> Any:
        """Reconstruct the canonical workbook payload from retained XLSX bytes."""
        limit = definition.maximum_workbook_bytes
        if limit is not None and len(raw_payload) > limit:
            raise AcquisitionError(
                "Retained XLSX artifact exceeds configured byte bound", code="BYTE_BOUND"
            )
        try:
            workbook = load_workbook(io.BytesIO(raw_payload), read_only=True, data_only=True)
        except (OSError, ValueError, KeyError) as error:
            raise AcquisitionError(
                "Retained XLSX artifact was not a readable workbook", code="WORKBOOK_INVALID"
            ) from error
        try:
            if definition.workbook_sheet not in workbook.sheetnames:
                raise AcquisitionError(
                    "Configured workbook sheet was not found", code="SHEET_NOT_FOUND"
                )
            sheet = workbook[definition.workbook_sheet]
            header_row = definition.header_row or 1
            rows_iter = sheet.iter_rows(values_only=True)
            headers_row: tuple[Any, ...] | None = None
            for index, row in enumerate(rows_iter, start=1):
                if index == header_row:
                    headers_row = tuple(row)
                    break
            if headers_row is None:
                raise AcquisitionError(
                    "Configured workbook header row was not found", code="HEADER_NOT_FOUND"
                )
            names = [str(value).strip() for value in headers_row]
            if len(names) != len(set(names)) or any(not name for name in names):
                raise AcquisitionError(
                    "Workbook header is not unique and non-empty", code="HEADER_INVALID"
                )
            missing = [column for column in definition.required_columns if column not in names]
            if missing:
                raise AcquisitionError(
                    f"Workbook is missing required columns: {', '.join(missing)}",
                    code="MISSING_COLUMNS",
                )
            rows: list[dict[str, Any]] = []
            for values in rows_iter:
                if not any(value is not None for value in values):
                    continue
                rows.append({name: value for name, value in zip(names, values, strict=False)})
                if definition.maximum_rows is not None and len(rows) > definition.maximum_rows:
                    raise AcquisitionError(
                        "Workbook exceeds configured row bound", code="ROW_BOUND"
                    )
            return {"sample": rows, "workbook_sheet": definition.workbook_sheet, "headers": names}
        finally:
            workbook.close()

    def normalize(self, definition: SourceDefinition, payload: Any) -> NormalizeResult:
        sample = payload.get("sample") if isinstance(payload, dict) else payload
        assert isinstance(sample, list)
        records = [_normalized_record(definition, row) for row in sample if isinstance(row, dict)]
        return NormalizeResult(
            records=records,
            transformation_version="http_xlsx_normalize_v1",
            detail={"record_count": len(records)},
        )


_REGISTRY: dict[AdapterKind, SourceAdapter] = {
    AdapterKind.SOCRATA: SocrataAdapter(),
    AdapterKind.HTTP_XLSX: HttpXlsxAdapter(),
}


def get_adapter(kind: AdapterKind) -> SourceAdapter:
    try:
        return _REGISTRY[kind]
    except KeyError as error:
        raise ValueError(
            f"Unsupported adapter_kind={kind!s}. Extend the registry with an explicit "
            "adapter module rather than encoding source logic in workflow YAML."
        ) from error


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _normalized_record(definition: SourceDefinition, row: dict[str, Any]) -> dict[str, Any]:
    """Add lineage without inventing a key for the historical CDC grain."""
    result: dict[str, Any] = {
        "source_id": definition.source_id,
        "dataset_id": definition.dataset_id,
        "source_definition_version": definition.definition_version,
        "geography_semantics": definition.geography_semantics,
        "temporal_semantics": definition.temporal_semantics,
        "record": row,
    }
    if definition.extra.get("normalization_kind") == "cdc_historical_2008_2021":
        result.update(
            {
                "surveillance_era": "2008-2021",
                "source_resolution": "COUNTY_YEAR_CASE_STATUS_SEX_AGE",
                "temporal_window": "ANNUAL_SURVEILLANCE_YEAR",
                "caveat": (
                    "2008-2021 surveillance era; county of residence, not exposure; "
                    "reported cases are not infection incidence; no direct comparison "
                    "with 2022 onward without reviewed methodology"
                ),
            }
        )
    return result


def _validate_era(
    definition: SourceDefinition,
    sample: list[Any],
    issues: list[ValidationIssue],
) -> None:
    minimum = definition.extra.get("minimum_year")
    maximum = definition.extra.get("maximum_year")
    if minimum is None or maximum is None:
        return
    for row in sample:
        if not isinstance(row, dict) or row.get("year") is None:
            continue
        try:
            year = int(str(row["year"]))
        except (TypeError, ValueError):
            issues.append(
                ValidationIssue(
                    code="YEAR_INVALID",
                    message="Source year is not an integer",
                    category=FailureCategory.SCHEMA,
                )
            )
            continue
        if not int(minimum) <= year <= int(maximum):
            issues.append(
                ValidationIssue(
                    code="YEAR_OUT_OF_RANGE",
                    message=f"Source year must be between {minimum} and {maximum}",
                    category=FailureCategory.SCHEMA,
                )
            )
