"""Evidence-only CDC Ixodes county-status onboarding for isolated DEV."""

from __future__ import annotations

import hashlib
import json
import os
import re
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
import yaml  # type: ignore[import-untyped]
from lyme_gap_atlas_shared.settings import SnowflakeSettings
from lyme_gap_atlas_shared.snowflake import connect
from openpyxl import load_workbook  # type: ignore[import-untyped]

from .artifacts import Artifact, create_artifact
from .assessment import Assessment
from .cdc import _spaces_client
from .redaction import redact_mapping
from .settings import PipelineSettings

RESOURCE_KEY = "cdc_tick_ixodes_county_status"
SOURCE_DATASET_ID = "cdc-ixodes-county-status-2025"
PROFILE_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "sources" / "cdc_tick_ixodes_county_status.yml"
)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BROWSER_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36"
)
EXPECTED_HEADERS = (
    "FIPSCode",
    "State",
    "County",
    "Ixodes_scapularis_County_Status",
    "Ixodes_scapularis_data_source",
    "Ixodes_pacificus_county_status",
    "Ixodes_pacificus_data_source",
)


@dataclass(frozen=True)
class FetchResult:
    payload: bytes
    media_type: str
    etag: str | None
    last_modified: str | None


@dataclass(frozen=True)
class WorkbookEvidence:
    sample: list[dict[str, object]]
    row_count: int
    schema: dict[str, object]


@dataclass(frozen=True)
class AcquisitionBundle:
    landing: FetchResult
    workbook: FetchResult
    manifest: dict[str, object]
    manifest_payload: bytes


EVIDENCE_ROUTE = "GITHUB_ACTIONS_CDC_EVIDENCE_V1"
OPERATOR_EVIDENCE_ROUTE = "OPERATOR_BROWSER_EXPORT_V1"
PDF_MEDIA_TYPE = "application/pdf"
IMAGE_DIGEST_PATTERN = re.compile(r"sha256:[0-9a-f]{64}")


def _manifest_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("CDC evidence manifest contains a duplicate key")
        result[key] = value
    return result


def load_tick_profile(path: Path = PROFILE_PATH) -> dict[str, Any]:
    """Load the pinned first-party CDC workbook evidence profile."""
    profile = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(profile, dict):
        raise ValueError("Invalid tick-surveillance source profile")
    expected = {
        "resource_key": RESOURCE_KEY,
        "source_dataset_id": SOURCE_DATASET_ID,
        "connector_name": "HTTP_XLSX_V1",
        "landing_page_url": "https://www.cdc.gov/ticks/data-research/facts-stats/tick-surveillance-data-sets.html",
        "endpoint_template": "https://www.cdc.gov/ticks/media/files/2026/04/Public_Use_Ixodes_County_Table_2026_03252026.xlsx",
        "workbook_sheet": "Ixodes records 2025",
        "header_row": 2,
        "deterministic_order_clause": "FIPSCode ASC",
        "incremental_strategy": "SNAPSHOT_DIFF",
    }
    if any(profile.get(key) != value for key, value in expected.items()):
        raise ValueError("Tick-surveillance source profile does not match the reviewed CDC source")
    if profile.get("onboarding_environments") != ["dev"]:
        raise ValueError("Tick-surveillance evidence onboarding must remain DEV-only")
    return profile


def _fetch_bytes(
    url: str,
    *,
    accept: str,
    maximum_bytes: int,
    referer: str | None = None,
    retries: int = 3,
) -> FetchResult:
    """Fetch one exact public CDC resource with a strict response bound."""
    if urlparse(url).scheme != "https" or urlparse(url).hostname != "www.cdc.gov":
        raise ValueError("Tick-surveillance evidence allows only first-party CDC HTTPS resources")
    headers = {
        "Accept": accept,
        "Accept-Language": "en-US,en;q=0.9",
        "Cache-Control": "no-cache",
        "Pragma": "no-cache",
        "User-Agent": BROWSER_USER_AGENT,
    }
    if referer:
        headers["Referer"] = referer
    for attempt in range(retries):
        try:
            with (
                httpx.Client(follow_redirects=True, timeout=30) as client,
                client.stream("GET", url, headers=headers) as response,
            ):
                if response.status_code == 429 or response.status_code >= 500:
                    response.raise_for_status()
                response.raise_for_status()
                final = urlparse(str(response.url))
                if final.scheme != "https" or final.hostname != "www.cdc.gov":
                    raise ValueError(
                        "CDC evidence request redirected outside the approved publisher"
                    )
                declared = response.headers.get("content-length")
                if declared and int(declared) > maximum_bytes:
                    raise ValueError("CDC evidence response exceeds the configured byte bound")
                payload = bytearray()
                for chunk in response.iter_bytes():
                    payload.extend(chunk)
                    if len(payload) > maximum_bytes:
                        raise ValueError("CDC evidence response exceeds the configured byte bound")
                media_type = response.headers.get("content-type", "").split(";", 1)[0].strip()
                return FetchResult(
                    bytes(payload),
                    media_type,
                    response.headers.get("etag"),
                    response.headers.get("last-modified"),
                )
        except httpx.HTTPStatusError as error:
            if (
                error.response.status_code not in {429, 500, 502, 503, 504}
                or attempt == retries - 1
            ):
                raise
            time.sleep(2**attempt)
    raise AssertionError("unreachable")


def _failure_classification(error: Exception) -> str:
    """Return a bounded, non-secret classification for retained failure evidence."""
    if isinstance(error, httpx.HTTPStatusError):
        return f"HTTP_{error.response.status_code}"
    if isinstance(error, ValueError):
        return "SOURCE_VALIDATION_FAILED"
    return "EVIDENCE_CAPTURE_FAILED"


def _validate_landing_page(result: FetchResult) -> None:
    if result.media_type != "text/html":
        raise ValueError("CDC tick-surveillance landing page did not return HTML")
    text = result.payload.decode("utf-8", errors="replace").lower()
    required = (
        "tick surveillance data sets",
        "public_use_ixodes_county_table_2026_03252026.xlsx",
        "no records",
        "established",
    )
    if any(value not in text for value in required):
        raise ValueError("CDC tick-surveillance landing-page semantics changed")


def _validate_landing_page_pdf(result: FetchResult) -> None:
    """Validate an opaque operator print without executing or rendering it."""
    if result.media_type != PDF_MEDIA_TYPE:
        raise ValueError("CDC tick-surveillance landing-page print is not a PDF")
    if not result.payload.startswith(b"%PDF-") or b"%%EOF" not in result.payload[-4096:]:
        raise ValueError("CDC tick-surveillance landing-page print is malformed")


def _load_evidence_bundle(bundle_dir: Path, profile: dict[str, Any]) -> AcquisitionBundle:
    """Load and independently verify the sealed source-evidence envelope."""
    manifest_path = bundle_dir / "acquisition-manifest.json"
    manifest_payload = manifest_path.read_bytes()
    if len(manifest_payload) > 100_000:
        raise ValueError("CDC evidence manifest exceeds the configured byte bound")
    manifest = json.loads(manifest_payload, object_pairs_hook=_manifest_object)
    if not isinstance(manifest, dict):
        raise ValueError("CDC evidence manifest must be a JSON object")
    route = manifest.get("acquisition_route")
    version = manifest.get("manifest_version")
    if (version, route) not in {
        (1, EVIDENCE_ROUTE),
        (2, OPERATOR_EVIDENCE_ROUTE),
    }:
        raise ValueError("CDC evidence manifest version or acquisition route is not approved")
    provenance_key = "github" if route == EVIDENCE_ROUTE else "operator"
    if set(manifest) != {
        "manifest_version",
        "acquisition_route",
        "retrieved_at",
        provenance_key,
        "base_image_digest",
        "resources",
    }:
        raise ValueError("CDC evidence manifest fields do not match the approved contract")
    retrieved_at = manifest.get("retrieved_at")
    if not isinstance(retrieved_at, str):
        raise ValueError("CDC evidence manifest is missing its retrieval timestamp")
    try:
        parsed_retrieved_at = datetime.fromisoformat(retrieved_at.replace("Z", "+00:00"))
    except ValueError as error:
        raise ValueError("CDC evidence manifest retrieval timestamp is invalid") from error
    if parsed_retrieved_at.tzinfo is None:
        raise ValueError("CDC evidence manifest retrieval timestamp must include a timezone")
    age = datetime.now(UTC) - parsed_retrieved_at.astimezone(UTC)
    maximum_age_seconds = 21_600 if route == EVIDENCE_ROUTE else 259_200
    if age.total_seconds() < -300 or age.total_seconds() > maximum_age_seconds:
        raise ValueError("CDC evidence manifest retrieval timestamp is outside the run window")

    github: dict[str, object] | None = None
    operator: dict[str, object] | None = None
    if route == EVIDENCE_ROUTE:
        github_value = manifest.get("github")
        if not isinstance(github_value, dict) or set(github_value) != {
            "repository",
            "run_id",
            "run_attempt",
            "sha",
        }:
            raise ValueError("CDC evidence manifest GitHub provenance is incomplete")
        github = github_value
        if github.get("repository") != "Caraway-Labs/one-health-lyme-gap-atlas-data":
            raise ValueError("CDC evidence manifest repository is not approved")
        if not all(isinstance(github.get(key), str) and github[key] for key in github):
            raise ValueError("CDC evidence manifest GitHub provenance values are invalid")
        if re.fullmatch(r"[0-9]+", str(github["run_id"])) is None:
            raise ValueError("CDC evidence manifest GitHub run ID is invalid")
        if re.fullmatch(r"[0-9]+", str(github["run_attempt"])) is None:
            raise ValueError("CDC evidence manifest GitHub run attempt is invalid")
        if re.fullmatch(r"[0-9a-f]{40}", str(github["sha"])) is None:
            raise ValueError("CDC evidence manifest GitHub commit is invalid")
    else:
        operator_value = manifest.get("operator")
        if not isinstance(operator_value, dict) or set(operator_value) != {
            "retrieval_id",
            "acquisition_method",
            "attestation",
        }:
            raise ValueError("CDC evidence manifest operator provenance is incomplete")
        operator = operator_value
        try:
            uuid.UUID(str(operator["retrieval_id"]))
        except (ValueError, TypeError) as error:
            raise ValueError("CDC evidence operator retrieval ID is invalid") from error
        if (
            operator.get("acquisition_method") != "BROWSER_DOWNLOAD_AND_PRINT"
            or operator.get("attestation") != "FILES_SAVED_FROM_PINNED_FIRST_PARTY_CDC_PAGE"
        ):
            raise ValueError("CDC evidence operator acquisition attestation is not approved")

    base_digest = manifest.get("base_image_digest")
    envelope_digest = os.getenv("TICK_EVIDENCE_ENVELOPE_DIGEST", "")
    runtime_base_digest = os.getenv("TICK_EVIDENCE_BASE_IMAGE_DIGEST", "")
    runtime_run_id = os.getenv("TICK_EVIDENCE_GITHUB_RUN_ID", "")
    runtime_retrieval_id = os.getenv("TICK_EVIDENCE_OPERATOR_RETRIEVAL_ID", "")
    if (
        not isinstance(base_digest, str)
        or IMAGE_DIGEST_PATTERN.fullmatch(base_digest) is None
        or runtime_base_digest != base_digest
    ):
        raise ValueError("CDC evidence base image digest is missing or does not match")
    if IMAGE_DIGEST_PATTERN.fullmatch(envelope_digest) is None:
        raise ValueError("CDC evidence envelope image digest is missing or invalid")
    if route == EVIDENCE_ROUTE and github is not None and runtime_run_id != github["run_id"]:
        raise ValueError("CDC evidence GitHub run provenance does not match the runtime")
    if (
        route == OPERATOR_EVIDENCE_ROUTE
        and operator is not None
        and runtime_retrieval_id != operator["retrieval_id"]
    ):
        raise ValueError("CDC evidence operator provenance does not match the runtime")

    landing_purpose = (
        "SOURCE_LANDING_PAGE" if route == EVIDENCE_ROUTE else "SOURCE_LANDING_PAGE_PRINT"
    )
    expected = {
        landing_purpose: (
            "landing.html" if route == EVIDENCE_ROUTE else "landing.pdf",
            str(profile["landing_page_url"]),
            "text/html" if route == EVIDENCE_ROUTE else PDF_MEDIA_TYPE,
            2_000_000,
        ),
        "SOURCE_WORKBOOK_EVIDENCE": (
            "workbook.xlsx",
            str(profile["endpoint_template"]),
            XLSX_MEDIA_TYPE,
            int(profile["maximum_workbook_bytes"]),
        ),
    }
    resources = manifest.get("resources")
    if not isinstance(resources, list) or len(resources) != len(expected):
        raise ValueError("CDC evidence manifest must contain exactly two resources")
    verified: dict[str, FetchResult] = {}
    for resource in resources:
        if not isinstance(resource, dict):
            raise ValueError("CDC evidence manifest resource is invalid")
        required_resource_fields = {
            "purpose",
            "filename",
            "requested_url",
            "final_url",
            "status_code",
            "media_type",
            "byte_count",
            "sha256",
            "etag",
            "last_modified",
        }
        if route == OPERATOR_EVIDENCE_ROUTE:
            required_resource_fields |= {
                "transport",
                "http_status_observed",
                "source_file_modified_at",
            }
        if set(resource) != required_resource_fields:
            raise ValueError("CDC evidence manifest resource fields do not match the contract")
        purpose = resource.get("purpose")
        if not isinstance(purpose, str) or purpose not in expected or purpose in verified:
            raise ValueError("CDC evidence manifest resource purpose is invalid or duplicated")
        filename, requested_url, media_type, maximum_bytes = expected[purpose]
        if resource.get("filename") != filename or resource.get("requested_url") != requested_url:
            raise ValueError("CDC evidence manifest resource identity changed")
        final_url = resource.get("final_url")
        if not isinstance(final_url, str):
            raise ValueError("CDC evidence manifest final URL is invalid")
        parsed_final = urlparse(final_url)
        if parsed_final.scheme != "https" or parsed_final.hostname != "www.cdc.gov":
            raise ValueError("CDC evidence manifest redirected outside the approved publisher")
        if route == EVIDENCE_ROUTE:
            if resource.get("status_code") != 200 or resource.get("media_type") != media_type:
                raise ValueError("CDC evidence manifest response status or media type changed")
        else:
            expected_transport = (
                "BROWSER_PRINT_TO_PDF"
                if purpose == "SOURCE_LANDING_PAGE_PRINT"
                else "BROWSER_DOWNLOAD"
            )
            if (
                resource.get("status_code") is not None
                or resource.get("http_status_observed") is not False
                or resource.get("media_type") != media_type
                or resource.get("transport") != expected_transport
            ):
                raise ValueError("CDC operator evidence transport metadata is invalid")
            modified_at = resource.get("source_file_modified_at")
            if not isinstance(modified_at, str):
                raise ValueError("CDC operator evidence file timestamp is invalid")
            try:
                parsed_modified_at = datetime.fromisoformat(modified_at.replace("Z", "+00:00"))
            except ValueError as error:
                raise ValueError("CDC operator evidence file timestamp is invalid") from error
            if (
                parsed_modified_at.tzinfo is None
                or abs((parsed_retrieved_at - parsed_modified_at).total_seconds()) > 3600
            ):
                raise ValueError("CDC operator evidence file timestamp is outside the capture")
        payload = (bundle_dir / filename).read_bytes()
        if not payload or len(payload) > maximum_bytes:
            raise ValueError("CDC evidence bundle resource is outside its byte bound")
        digest = hashlib.sha256(payload).hexdigest()
        if resource.get("byte_count") != len(payload) or resource.get("sha256") != digest:
            raise ValueError("CDC evidence bundle checksum or byte count does not match")
        etag = resource.get("etag")
        last_modified = resource.get("last_modified")
        if etag is not None and not isinstance(etag, str):
            raise ValueError("CDC evidence manifest ETag is invalid")
        if last_modified is not None and not isinstance(last_modified, str):
            raise ValueError("CDC evidence manifest Last-Modified value is invalid")
        verified[purpose] = FetchResult(payload, media_type, etag, last_modified)
    return AcquisitionBundle(
        landing=verified[landing_purpose],
        workbook=verified["SOURCE_WORKBOOK_EVIDENCE"],
        manifest=manifest,
        manifest_payload=manifest_payload,
    )


def _parse_workbook(payload: bytes, profile: dict[str, Any], sample_limit: int) -> WorkbookEvidence:
    if not zipfile.is_zipfile(BytesIO(payload)):
        raise ValueError("CDC tick-surveillance evidence is not a valid XLSX package")
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        if sum(member.file_size for member in archive.infolist()) > int(
            profile["maximum_uncompressed_bytes"]
        ):
            raise ValueError("CDC workbook exceeds the configured uncompressed byte bound")
        if any(
            member.filename.startswith(("/", "\\")) or ".." in Path(member.filename).parts
            for member in archive.infolist()
        ):
            raise ValueError("CDC workbook contains an invalid package member")
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    required_sheets = {
        "Data Use Agreement",
        "Classification Terms",
        str(profile["workbook_sheet"]),
    }
    if len(workbook.sheetnames) > 10 or not required_sheets.issubset(workbook.sheetnames):
        raise ValueError("CDC workbook sheet structure changed")
    agreement_text = " ".join(
        str(value)
        for row in workbook["Data Use Agreement"].iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()
    classification_text = " ".join(
        str(value)
        for row in workbook["Classification Terms"].iter_rows(values_only=True)
        for value in row
        if value is not None
    ).lower()
    required_agreement_terms = (
        "access to arbonet tick module data is limited",
        "should not be provided to other persons",
        "appropriately referenced",
        "final copy",
        "passive surveillance system",
    )
    required_classification_terms = (
        "established",
        "reported",
        "no records",
        "should not be interpreted as ticks being absent",
    )
    if any(term not in agreement_text for term in required_agreement_terms) or any(
        term not in classification_text for term in required_classification_terms
    ):
        raise ValueError("CDC workbook data-use or classification semantics changed")
    sheet = workbook[str(profile["workbook_sheet"])]
    if sheet.max_row is None or sheet.max_column is None or not 3 <= sheet.max_row <= 10_000:
        raise ValueError("CDC workbook row shape is outside the reviewed evidence bound")
    if sheet.max_column != len(EXPECTED_HEADERS):
        raise ValueError("CDC workbook column count changed")
    header_row = int(profile["header_row"])
    headers = tuple(
        str(value) if value is not None else ""
        for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    )
    if headers != EXPECTED_HEADERS:
        raise ValueError("CDC workbook schema changed; steward review is required")
    sample: list[dict[str, object]] = []
    prior_fips = ""
    row_count = 0
    allowed = set(str(value) for value in profile["allowed_status_values"])
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in values):
            continue
        row = dict(zip(headers, values, strict=True))
        fips = row["FIPSCode"]
        if not isinstance(fips, str) or re.fullmatch(r"[0-9]{5}", fips) is None:
            raise ValueError("CDC workbook sample does not preserve five-character county FIPS")
        if prior_fips and fips <= prior_fips:
            raise ValueError("CDC workbook sample is not deterministically ordered by county FIPS")
        prior_fips = fips
        for field in ("Ixodes_scapularis_County_Status", "Ixodes_pacificus_county_status"):
            if row[field] not in allowed:
                raise ValueError("CDC workbook contains an unreviewed county-status value")
        row_count += 1
        if len(sample) < sample_limit:
            sample.append(row)
    if len(sample) != sample_limit:
        raise ValueError("CDC workbook does not contain the requested bounded sample")
    schema = {
        "contract_version": "tick-surveillance-v1",
        "workbook_sheet": sheet.title,
        "header_row": header_row,
        "headers": list(headers),
        "worksheet_rows_including_headers": sheet.max_row,
        "worksheet_columns": sheet.max_column,
        "dataset_as_of": str(profile["dataset_as_of"]),
        "temporal_semantics": str(profile["temporal_semantics"]),
        "allowed_status_values": list(profile["allowed_status_values"]),
        "embedded_data_use_agreement_validated": True,
        "embedded_classification_terms_validated": True,
        "full_dataset_quality_validated": False,
    }
    return WorkbookEvidence(sample, row_count, schema)


def _save_artifact(
    s3: Any,
    settings: PipelineSettings,
    run_id: str,
    payload: bytes,
    media_type: str,
) -> Artifact:
    artifact = create_artifact(
        payload=payload,
        environment=settings.topx_env,
        resource_key=RESOURCE_KEY,
        run_id=run_id,
    )
    s3.put_object(
        Bucket=settings.spaces_bucket,
        Key=f"{settings.spaces_prefix}/{artifact.object_key}",
        Body=payload,
        ContentType=media_type,
    )
    return artifact


def collect_tick_surveillance_evidence(
    sample_limit: int = 25, *, evidence_bundle_dir: Path | None = None
) -> dict[str, object]:
    """Create a pending DEV candidate without RAW loading, approval, or transformation."""
    if not 1 <= sample_limit <= 100:
        raise ValueError("sample_limit must be between 1 and 100")
    settings = PipelineSettings()
    if settings.topx_env != "dev":
        raise ValueError("Tick-surveillance evidence capture is approved only for isolated DEV")
    profile = load_tick_profile()
    landing_url = str(profile["landing_page_url"])
    workbook_url = str(profile["endpoint_template"])
    run_id = str(uuid.uuid4())
    now = datetime.now(UTC)
    profile_sha256 = hashlib.sha256(
        yaml.safe_dump(profile, sort_keys=True).encode("utf-8")
    ).hexdigest()

    with connect(SnowflakeSettings()) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """INSERT INTO GOVERNANCE.INGESTION_RUNS
                (ingestion_run_id, resource_key, run_mode, trigger_type, status, code_version,
                 config_sha256, started_at)
                VALUES (%s, %s, 'EVIDENCE_ONLY', 'MANUAL', 'RUNNING',
                        'cdc-tick-ixodes-evidence-v1', %s, %s)""",
                (run_id, RESOURCE_KEY, profile_sha256, now),
            )
        connection.commit()
        try:
            bundle = (
                _load_evidence_bundle(evidence_bundle_dir, profile)
                if evidence_bundle_dir is not None
                else None
            )
            landing = (
                bundle.landing
                if bundle is not None
                else _fetch_bytes(landing_url, accept="text/html", maximum_bytes=2_000_000)
            )
            if landing.media_type == "text/html":
                _validate_landing_page(landing)
            else:
                _validate_landing_page_pdf(landing)
            workbook = (
                bundle.workbook
                if bundle is not None
                else _fetch_bytes(
                    workbook_url,
                    accept=XLSX_MEDIA_TYPE,
                    maximum_bytes=int(profile["maximum_workbook_bytes"]),
                    referer=landing_url,
                )
            )
            if workbook.media_type != XLSX_MEDIA_TYPE:
                raise ValueError("CDC tick-surveillance workbook media type changed")
            evidence = _parse_workbook(workbook.payload, profile, sample_limit)
            schema_payload = json.dumps(evidence.schema, sort_keys=True, separators=(",", ":"))
            schema_sha256 = hashlib.sha256(schema_payload.encode()).hexdigest()
            sample_payload = json.dumps(
                evidence.sample, sort_keys=True, separators=(",", ":")
            ).encode()
            metadata = {
                "title": "Blacklegged and western blacklegged tick county status",
                "publisher": "CDC National Center for Emerging and Zoonotic Infectious Diseases",
                "source_dataset_id": SOURCE_DATASET_ID,
                "landing_page_url": landing_url,
                "workbook_url": workbook_url,
                "dataset_as_of": str(profile["dataset_as_of"]),
                "workbook_etag": workbook.etag,
                "workbook_last_modified": workbook.last_modified,
                "workbook_rows": evidence.row_count,
                "sample_rows": len(evidence.sample),
                "county_status_semantics": {
                    "Established": "Publisher cumulative established classification",
                    "Reported": "Publisher cumulative reported classification",
                    "No records": "No reported surveillance evidence; not evidence of absence",
                },
                "license_or_terms_status": (
                    "RESTRICTED_REVIEW_REQUIRED_EMBEDDED_DATA_USE_AGREEMENT"
                ),
                "full_dataset_quality_validated": False,
                "acquisition_route": (
                    str(bundle.manifest["acquisition_route"])
                    if bundle is not None
                    else "DIGITALOCEAN_DIRECT_CDC_V1"
                ),
                "base_image_digest": (
                    str(bundle.manifest["base_image_digest"]) if bundle is not None else None
                ),
                "envelope_image_digest": (
                    os.getenv("TICK_EVIDENCE_ENVELOPE_DIGEST") if bundle is not None else None
                ),
                "github": (bundle.manifest.get("github") if bundle is not None else None),
                "operator": (bundle.manifest.get("operator") if bundle is not None else None),
            }
            metadata_payload = json.dumps(metadata, sort_keys=True, separators=(",", ":")).encode()
            s3 = _spaces_client(settings)
            landing_artifact = _save_artifact(
                s3, settings, run_id, landing.payload, landing.media_type
            )
            workbook_artifact = _save_artifact(
                s3, settings, run_id, workbook.payload, XLSX_MEDIA_TYPE
            )
            metadata_artifact = _save_artifact(
                s3, settings, run_id, metadata_payload, "application/json"
            )
            sample_artifact = _save_artifact(
                s3, settings, run_id, sample_payload, "application/json"
            )
            manifest_artifact = (
                _save_artifact(
                    s3,
                    settings,
                    run_id,
                    bundle.manifest_payload,
                    "application/json",
                )
                if bundle is not None
                else None
            )
            assessment = Assessment(95, 95, 65, 90, 65)
            limitations = (
                "Cumulative county status through 2025-12-31; No records is not tick absence. "
                "The workbook does not provide collection effort, abundance, life stage, or "
                "pathogen testing. Evidence capture retained the byte-bounded publisher workbook "
                "because no row API exists. It validated schema, keys, ordering, and status "
                "domains across the worksheet but serialized only a bounded sample; it did not "
                "load RAW or "
                "run the complete post-ingestion quality suite. The embedded agreement and source "
                "citations restrict raw redistribution, require ArboNET attribution, and require "
                "delivery of a final publication copy to CDC; the steward must review these terms."
            )
            artifact_ids = {
                name: str(uuid.uuid4())
                for name in ("landing", "workbook", "metadata", "sample", "manifest")
            }
            connection.autocommit(False)
            with connection.cursor() as cursor:
                catalog_dataset_id = str(uuid.uuid4())
                cursor.execute(
                    """INSERT INTO GOVERNANCE.CATALOG_DATASETS
                    (catalog_dataset_id, dataset_key, catalog_name, catalog_record_id,
                     metadata_payload, metadata_sha256, discovered_at, is_current)
                    SELECT %s, %s, 'CDC_WEB', %s, PARSE_JSON(%s), %s, %s, TRUE""",
                    (
                        catalog_dataset_id,
                        RESOURCE_KEY,
                        SOURCE_DATASET_ID,
                        json.dumps(metadata),
                        metadata_artifact.sha256,
                        now,
                    ),
                )
                cursor.execute(
                    """INSERT INTO GOVERNANCE.CATALOG_RESOURCES
                    (catalog_resource_id, catalog_dataset_id, resource_key, resource_type,
                     resource_url, canonical_source_url, api_dataset_id, resource_payload,
                     registered_at, is_active)
                    SELECT %s, %s, %s, 'DATA', %s, %s, %s, PARSE_JSON(%s), %s, TRUE""",
                    (
                        str(uuid.uuid4()),
                        catalog_dataset_id,
                        RESOURCE_KEY,
                        workbook_url,
                        landing_url,
                        SOURCE_DATASET_ID,
                        json.dumps(
                            {
                                "title": metadata["title"],
                                "publisher": metadata["publisher"],
                                "media_type": XLSX_MEDIA_TYPE,
                                "terms_status": metadata["license_or_terms_status"],
                            }
                        ),
                        now,
                    ),
                )
                cursor.execute(
                    "UPDATE GOVERNANCE.SOURCE_ACCESS_PROFILES SET effective_to=%s "
                    "WHERE resource_key=%s AND effective_to IS NULL",
                    (now, RESOURCE_KEY),
                )
                cursor.execute(
                    """INSERT INTO GOVERNANCE.SOURCE_ACCESS_PROFILES
                    (source_access_profile_id, resource_key, profile_version, connector_name,
                     endpoint_template, deterministic_order_clause, incremental_strategy,
                     configuration_sha256, effective_from)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (
                        str(uuid.uuid4()),
                        RESOURCE_KEY,
                        int(profile["profile_version"]),
                        str(profile["connector_name"]),
                        workbook_url,
                        str(profile["deterministic_order_clause"]),
                        str(profile["incremental_strategy"]),
                        profile_sha256,
                        now,
                    ),
                )
                operator_route = metadata["acquisition_route"] == OPERATOR_EVIDENCE_ROUTE
                landing_purpose = (
                    "SOURCE_LANDING_PAGE_PRINT" if operator_route else "SOURCE_LANDING_PAGE"
                )
                request_status = None if operator_route else 200
                requests = (
                    (
                        1,
                        landing_purpose,
                        landing_url,
                        landing_artifact,
                        1,
                        request_status,
                    ),
                    (
                        2,
                        "SOURCE_WORKBOOK_EVIDENCE",
                        workbook_url,
                        workbook_artifact,
                        evidence.row_count,
                        request_status,
                    ),
                )
                request_ids: dict[str, str] = {}
                for sequence, purpose, endpoint, artifact, row_count, status_code in requests:
                    request_id = str(uuid.uuid4())
                    request_ids[purpose] = request_id
                    cursor.execute(
                        """INSERT INTO GOVERNANCE.INGESTION_REQUESTS
                        (ingestion_request_id, ingestion_run_id, request_sequence, request_purpose,
                         endpoint, redacted_request, status_code, response_sha256,
                         retrieved_row_count, created_at)
                        SELECT %s, %s, %s, %s, %s, PARSE_JSON(%s), %s, %s, %s, %s""",
                        (
                            request_id,
                            run_id,
                            sequence,
                            purpose,
                            endpoint,
                            json.dumps(
                                redact_mapping(
                                    {
                                        "endpoint": endpoint,
                                        "sample_limit": sample_limit,
                                        "acquisition_route": metadata["acquisition_route"],
                                        "base_image_digest": metadata["base_image_digest"],
                                        "envelope_image_digest": metadata["envelope_image_digest"],
                                        "github_run_id": (
                                            metadata["github"].get("run_id")
                                            if isinstance(metadata["github"], dict)
                                            else None
                                        ),
                                        "operator_retrieval_id": (
                                            metadata["operator"].get("retrieval_id")
                                            if isinstance(metadata["operator"], dict)
                                            else None
                                        ),
                                    }
                                )
                            ),
                            status_code,
                            artifact.sha256,
                            row_count,
                            now,
                        ),
                    )
                raw_artifacts: tuple[tuple[str, str, Artifact, str, str], ...] = (
                    (
                        "landing",
                        landing_purpose,
                        landing_artifact,
                        landing.media_type,
                        request_ids[landing_purpose],
                    ),
                    (
                        "workbook",
                        "SOURCE_WORKBOOK_EVIDENCE",
                        workbook_artifact,
                        XLSX_MEDIA_TYPE,
                        request_ids["SOURCE_WORKBOOK_EVIDENCE"],
                    ),
                    (
                        "metadata",
                        "NORMALIZED_SOURCE_METADATA",
                        metadata_artifact,
                        "application/json",
                        request_ids[landing_purpose],
                    ),
                    (
                        "sample",
                        "ORDERED_SOURCE_SAMPLE",
                        sample_artifact,
                        "application/json",
                        request_ids["SOURCE_WORKBOOK_EVIDENCE"],
                    ),
                )
                if manifest_artifact is not None:
                    raw_artifacts += (
                        (
                            "manifest",
                            "ACQUISITION_MANIFEST",
                            manifest_artifact,
                            "application/json",
                            request_ids[landing_purpose],
                        ),
                    )
                for name, artifact_type, artifact, media_type, request_id in raw_artifacts:
                    cursor.execute(
                        """INSERT INTO GOVERNANCE.RAW_ARTIFACTS
                        (artifact_id, ingestion_run_id, ingestion_request_id, artifact_uri,
                         artifact_type, media_type, byte_count, sha256, created_at)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                        (
                            artifact_ids[name],
                            run_id,
                            request_id,
                            f"s3://{settings.spaces_bucket}/{settings.spaces_prefix}/{artifact.object_key}",
                            artifact_type,
                            media_type,
                            artifact.byte_count,
                            artifact.sha256,
                            now,
                        ),
                    )
                for document_type, document_url, artifact_name, artifact in (
                    ("LANDING_PAGE_AND_METHODOLOGY", landing_url, "landing", landing_artifact),
                    (
                        "DATA_USE_AGREEMENT_AND_SOURCE_WORKBOOK",
                        workbook_url,
                        "workbook",
                        workbook_artifact,
                    ),
                ):
                    cursor.execute(
                        """INSERT INTO GOVERNANCE.SOURCE_DOCUMENT_SNAPSHOTS
                        (source_document_snapshot_id, resource_key, document_type, document_url,
                         artifact_id, content_sha256, retrieved_at, is_material_change)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, FALSE)""",
                        (
                            str(uuid.uuid4()),
                            RESOURCE_KEY,
                            document_type,
                            document_url,
                            artifact_ids[artifact_name],
                            artifact.sha256,
                            now,
                        ),
                    )
                cursor.execute(
                    """INSERT INTO GOVERNANCE.SCHEMA_SNAPSHOTS
                    (schema_snapshot_id, resource_key, schema_fingerprint, schema_payload,
                     retrieved_at)
                    SELECT %s, %s, %s, PARSE_JSON(%s), %s""",
                    (str(uuid.uuid4()), RESOURCE_KEY, schema_sha256, schema_payload, now),
                )
                cursor.execute(
                    """INSERT INTO GOVERNANCE.DATASET_QUALITY_ASSESSMENTS
                    (dataset_quality_assessment_id, resource_key, assessment_status,
                     relevance_score, joinability_score, accessibility_score, documentation_score,
                     quality_score, overall_score, recommendation, limitations, assessed_at)
                    VALUES (%s, %s, 'PENDING_REVIEW', 95, 95, 65, 90, 65, %s, %s, %s, %s)""",
                    (
                        str(uuid.uuid4()),
                        RESOURCE_KEY,
                        assessment.score,
                        assessment.recommendation,
                        limitations,
                        now,
                    ),
                )
                cursor.execute(
                    """UPDATE GOVERNANCE.INGESTION_RUNS
                    SET status='COMPLETED', completed_at=%s
                    WHERE ingestion_run_id=%s""",
                    (datetime.now(UTC), run_id),
                )
            connection.commit()
        except Exception as error:
            connection.rollback()
            with connection.cursor() as cursor:
                cursor.execute(
                    """UPDATE GOVERNANCE.INGESTION_RUNS
                    SET status='FAILED', completed_at=%s, error_classification=%s,
                        redacted_error=%s
                    WHERE ingestion_run_id=%s""",
                    (
                        datetime.now(UTC),
                        _failure_classification(error),
                        "CDC tick-surveillance evidence capture failed; review protected logs",
                        run_id,
                    ),
                )
            connection.commit()
            raise
    return {
        "ingestion_run_id": run_id,
        "resource_key": RESOURCE_KEY,
        "source_dataset_id": SOURCE_DATASET_ID,
        "sample_rows": len(evidence.sample),
        "workbook_rows": evidence.row_count,
        "landing_sha256": landing_artifact.sha256,
        "workbook_sha256": workbook_artifact.sha256,
        "sample_sha256": sample_artifact.sha256,
        "schema_sha256": schema_sha256,
        "acquisition_manifest_sha256": (
            manifest_artifact.sha256 if manifest_artifact is not None else None
        ),
        "full_dataset_quality_validated": False,
        "status": "PENDING_STEWARD_REVIEW",
    }
